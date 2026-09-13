from __future__ import annotations

import docx
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from web_portal.lib.docx_lang import set_doc_language_ru, set_run_language_ru


class InterceptsReport:
    """
    Экспорт Word для "Перехватов" в ТОМ ЖЕ формате, что и бланк из вкладки "Сеансы":
    таблица (Table Grid), Times New Roman 12, те же ширины колонок.

    Отличие:
      - вместо текста "В период ... отмечено ..." вставляем РЕАЛЬНОЕ содержание перехвата.
    """

    def __init__(self, *, position_name: str, started_at: str, ended_at: str | None):
        self.position_name = position_name
        self.started_at = started_at
        self.ended_at = ended_at
        self.rows_counter = 0
        self.doc = docx.Document()
        # Русский язык по умолчанию (для орфографии/Т9 в OnlyOffice)
        set_doc_language_ru(self.doc)

        header = ["", "Принадлежность", "Содержание перехвата", "Примечания"]
        self.table = self.doc.add_table(rows=1, cols=len(header))
        row = self.table.rows[0].cells
        for i, e in enumerate(header):
            row[i].text = e
            row[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT

    def _change_font(self):
        for row in self.table.rows:
            for cell in row.cells:
                if not cell.paragraphs:
                    continue
                # Обрабатываем все параграфы в ячейке
                for paragraph in cell.paragraphs:
                    # Устанавливаем язык и шрифт для всех runs в параграфе
                    if paragraph.runs:
                        for run in paragraph.runs:
                            font = run.font
                            font.name = "Times New Roman"
                            font.size = Pt(12)
                            set_run_language_ru(run)
                    elif paragraph.text and paragraph.text.strip():
                        # Если есть текст, но нет runs - создаём run
                        run = paragraph.add_run(paragraph.text)
                        paragraph.text = ""
                        font = run.font
                        font.name = "Times New Roman"
                        font.size = Pt(12)
                        set_run_language_ru(run)

    def _change_width(self):
        self.table.autofit = False
        self.table.allow_autofit = False  # type: ignore
        widths = (Cm(1), Cm(2.5), Cm(10.5), Cm(3))
        for col_idx, width in enumerate(widths):
            self.table.columns[col_idx].width = width
        for row in self.table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = width

    def add_item(
        self, idx: int, unit_name: str, frequency: str, group_code: str, content: str, 
        item_id: int = 0, location: str = ""
    ):
        # idx приходит снаружи (enumerate), но мы держим и свой счётчик — как в seanses
        self.rows_counter += 1
        row = self.table.add_row().cells
        row[0].text = f"{self.rows_counter}."

        # "Принадлежность" — формат: unit_name, frequency, ID, group_id, location
        parts = [
            str(unit_name or ""),
            str(frequency or ""),
            "ID",
            str(group_code or ""),
        ]
        if location and location.strip():
            parts.append(str(location or ""))
        row[1].text = "\n".join(parts)
        row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        # "Содержание перехвата" — реальный текст, сохраняем переносы строк.
        txt = str(content or "").rstrip("\n")
        cell = row[2]
        # В python-docx переносы в cell.text иногда ведут себя нестабильно в таблицах,
        # поэтому раскладываем текст по абзацам явно.
        lines = str(txt).replace("\r\n", "\n").replace("\r", "\n").split("\n")
        # гарантируем хотя бы один абзац
        if not cell.paragraphs:
            _p0 = cell.add_paragraph("")
        # очищаем первый абзац и заполняем
        cell.paragraphs[0].text = lines[0] if lines else ""
        cell.paragraphs[0].alignment = (
            WD_ALIGN_PARAGRAPH.LEFT if txt else WD_ALIGN_PARAGRAPH.CENTER
        )
        # добавляем оставшиеся строки как отдельные абзацы (сохраняет формат "в столбик")
        for ln in lines[1:] if len(lines) > 1 else []:
            p = cell.add_paragraph(ln)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT

        # Примечания — пока пусто (оставляем как поле под будущее)
        row[3].text = ""
        row[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    def save(self, path: str):
        self.table.style = "Table Grid"
        set_doc_language_ru(self.doc)
        self._change_font()
        self._change_width()
        self.doc.save(path)

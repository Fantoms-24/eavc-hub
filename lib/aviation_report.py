from __future__ import annotations

import docx
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from web_portal.lib.docx_lang import set_doc_language_ru, set_run_language_ru


class AviationReport:
    """
    Экспорт Word для "Авиации" в том же формате, что и бланк из вкладки "Сеансы":
    таблица (Table Grid), Times New Roman 12, те же ширины колонок.
    """

    def __init__(
        self,
        *,
        frequency: str,
        aviation_type: str,
        content: str,
        callsigns: list[str] | None = None,
    ):
        self.frequency = frequency
        self.aviation_type = aviation_type
        self.content = content
        self.callsigns = callsigns or []
        self.rows_counter = 0
        self.doc = docx.Document()
        # Русский язык по умолчанию (для орфографии/Т9 в OnlyOffice)
        set_doc_language_ru(self.doc)

        header = ["", "Принадлежность", "Содержание перехвата", "Примечания"]
        self.table = self.doc.add_table(rows=1, cols=len(header))
        row = self.table.rows[0].cells
        for i, e in enumerate(header):
            row[i].text = e
            row[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT

    def _change_font(self):
        for row in self.table.rows:
            for cell in row.cells:
                if not cell.paragraphs or not cell.paragraphs[0].runs:
                    continue
                run = cell.paragraphs[0].runs[0]
                font = run.font
                font.name = "Times New Roman"
                font.size = Pt(12)
                set_run_language_ru(run)

    def _change_width(self):
        self.table.autofit = False
        self.table.allow_autofit = False  # type: ignore
        widths = (Cm(1), Cm(2.5), Cm(8), Cm(3))
        for col_idx, width in enumerate(widths):
            self.table.columns[col_idx].width = width
        for row in self.table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = width

    def save(self, path: str):
        self.rows_counter += 1
        row = self.table.add_row().cells
        row[0].text = f"{self.rows_counter}."

        # "Принадлежность" — тип авиации и частота (без ID)
        row[1].text = f"{self.aviation_type}\n{self.frequency}"
        row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        # "Содержание перехвата" — реальный текст бланка
        txt = str(self.content or "").strip()
        row[2].text = txt
        row[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if txt else WD_ALIGN_PARAGRAPH.CENTER

        # Примечания — позывные, привязанные к этой частоте
        notes = "\n".join(self.callsigns) if self.callsigns else ""
        row[3].text = notes
        row[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        self.table.style = "Table Grid"
        set_doc_language_ru(self.doc)
        self._change_font()
        self._change_width()
        self.doc.save(path)

from __future__ import annotations

import io
import re
import tempfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from web_portal.config import BASE_DIR
from web_portal.lib.db import canonical_intensity_unit_key

TIME_SLOT_RE = re.compile(r"^\d{2}[.:]\d{2}-\d{2}[.:]\d{2}$", re.IGNORECASE)
CORR_LABEL_RE = re.compile(r"корреспонд", re.IGNORECASE)
UNIT_TYPE_RE = re.compile(
    r"(\d+)\s*(обрмп|обмп|minбатр|503|пдб|мб|шб|бон|батр|гадн|дшб|бмп|аэмб|брон|спб|бгтр|бтгр|ошб|"
    r"зрадн|мпб|баг|тб|тр|шспб|об(?!м|р)|еб(?!р)|обр(?!м)|бр(?!а))"
    r"\s*/?\s*(\d+)?",
    re.IGNORECASE,
)
LEADING_TYPE_RE = re.compile(
    r"^(тб|тр|шспб|зрадн|мпб|обмп|бмп|аэмб)\s+(\d+)",
    re.IGNORECASE,
)
QUOTE_RE = re.compile(r'["«»„“]')
NON_WORD_RE = re.compile(r"[^\w\d/]+", re.UNICODE)


def normalize_time_label(value: Any) -> str | None:
    text = str(value or "").strip()
    if not text:
        return None
    text = text.replace("–", "-").replace("—", "-").replace(":", ".")
    m = TIME_SLOT_RE.match(text)
    if not m:
        return None
    return text.replace(":", ".")


def _normalize_unit_text(name: str) -> str:
    n = str(name or "").replace("\n", " ")
    n = QUOTE_RE.sub("", n)
    n = re.sub(r"\b(\d+)/(\d+)\s+обрмп\b", r"\1 бмп \2 обрмп", n, flags=re.IGNORECASE)
    n = n.replace("/", " ")
    n = NON_WORD_RE.sub(" ", n)
    n = " ".join(n.split()).casefold()
    n = re.sub(r"\bбтгр\b", "бгтр", n)
    n = re.sub(r"\b(\d+)\s+еб\s+152\b", r"\1 бгтр 152", n)
    return n


def _unit_signature(name: str) -> tuple[str, ...]:
    n = _normalize_unit_text(name)

    obmp = re.match(r"^(\d+)\s+обмп\s+(\d+)\s+обрмп", n)
    if obmp:
        return (obmp.group(1), "обмп", obmp.group(2), "обрмп")

    bmp = re.match(r"^(\d+)\s+бмп\s+(\d+)\s+обрмп", n)
    if bmp:
        return (bmp.group(1), "бмп", bmp.group(2))

    lead = LEADING_TYPE_RE.match(n)
    if lead:
        return (lead.group(1).casefold(), lead.group(2))
    tokens: list[str] = []
    for m in UNIT_TYPE_RE.finditer(n):
        num1 = m.group(1)
        kind = m.group(2).casefold()
        num2 = (m.group(3) or "").strip()
        tokens.extend(x for x in (num1, kind, num2) if x)
    if tokens:
        return tuple(tokens)
    parts = [p for p in n.split() if p and p not in {"шв", "ак", "нгу", "кбр", "дшв", "сбпс", "кмп", "ошп"}]
    return tuple(parts[:6])


def _unit_match_score(target: str, source: str) -> float:
    ta = _normalize_unit_text(target)
    tb = _normalize_unit_text(source)
    if not ta or not tb:
        return 0.0
    if ta == tb:
        return 100.0

    if "шспб" in ta and "шспб" in tb:
        nums_t = set(re.findall(r"\d+", ta))
        nums_s = set(re.findall(r"\d+", tb))
        overlap = len(nums_t & nums_s)
        if overlap:
            return 78.0 + overlap * 3.0

    if "батар" in tb and "батар" not in ta:
        return 0.0

    sig_a = _unit_signature(target)
    sig_b = _unit_signature(source)
    if sig_a and sig_b:
        if sig_a == sig_b:
            anchor = " ".join(sig_a[: min(3, len(sig_a))])
            pos = tb.find(anchor) if anchor else -1
            if pos < 0:
                pos = ta.find(anchor) if anchor else -1
            if pos < 0:
                return 92.0
            return 95.0 - pos / 5.0 - len(tb) / 40.0
        short, long = (sig_a, sig_b) if len(sig_a) <= len(sig_b) else (sig_b, sig_a)
        if len(long) >= len(short) and long[: len(short)] == short:
            anchor = " ".join(short[: min(3, len(short))])
            pos = tb.find(anchor) if anchor else -1
            if pos < 0:
                pos = 999
            return 80.0 - pos / 5.0 - len(tb) / 40.0

    if ta in tb or tb in ta:
        return 60.0 + min(len(ta), len(tb)) / 10.0
    key_a = canonical_intensity_unit_key(target)
    key_b = canonical_intensity_unit_key(source)
    if key_a and key_b and key_a == key_b:
        return 55.0
    overlap = len(set(ta.split()) & set(tb.split()))
    if overlap >= 2:
        return 20.0 + overlap
    return 0.0


def match_units(target_units: list[str], source_units: list[str]) -> list[str | None]:
    src_unique = []
    seen: set[str] = set()
    for name in source_units:
        n = str(name or "").strip()
        if not n:
            continue
        key = n.casefold()
        if key in seen:
            continue
        seen.add(key)
        src_unique.append(n)

    out: list[str | None] = []
    for target in target_units:
        best_name: str | None = None
        best_score = 0.0
        for source in src_unique:
            score = _unit_match_score(target, source)
            if score > best_score:
                best_score = score
                best_name = source
        out.append(best_name if best_score >= 40.0 else None)
    return out


def _safe_int(value: Any) -> int:
    if value is None or value == "":
        return 0
    try:
        return int(round(float(value)))
    except (TypeError, ValueError):
        return 0


def _read_header_units(ws, header_row: int) -> list[str]:
    units: list[str] = []
    col = 2
    while col <= int(ws.max_column or 1):
        val = ws.cell(header_row, col).value
        if val is None or str(val).strip() == "":
            if units:
                break
            col += 1
            continue
        units.append(str(val).replace("\n", " ").strip())
        col += 1
    return units


def _read_current_period_block(ws, units: list[str]) -> tuple[str | None, dict[str, int], dict[str, int]]:
    """Текущий интервал экспорта портала — строки 5 (сеансы) и 6 (корреспонденты)."""
    current_slot = normalize_time_label(ws.cell(5, 1).value)
    if not current_slot:
        return None, {}, {}
    sessions: dict[str, int] = {}
    corr: dict[str, int] = {}
    for idx, unit in enumerate(units):
        col = idx + 2
        sessions[unit] = _safe_int(ws.cell(5, col).value)
        corr[unit] = _safe_int(ws.cell(6, col).value)
    return current_slot, sessions, corr


def parse_2h_workbook(content: bytes, *, current_period_only: bool = False) -> dict[str, Any]:
    wb = load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active
    header_row = 1
    label_a = str(ws.cell(1, 1).value or "").strip().casefold()
    if label_a not in {"сеть", "бланк"}:
        for r in range(1, min(6, int(ws.max_row or 1) + 1)):
            if str(ws.cell(r, 1).value or "").strip().casefold() in {"сеть", "бланк"}:
                header_row = r
                break

    units = _read_header_units(ws, header_row)
    if not units:
        wb.close()
        raise ValueError("В двухчасовой таблице не найдены подразделения")

    if current_period_only and header_row == 1:
        current_slot, sessions, corr = _read_current_period_block(ws, units)
        wb.close()
        if not current_slot:
            raise ValueError(
                "В двухчасовой таблице не найден текущий интервал (строка 5, формат 18.00-20.00)"
            )
        return {
            "units": units,
            "slots": {current_slot: {"sessions": sessions, "correspondents": corr}},
        }

    slots: dict[str, dict[str, dict[str, int]]] = {}

    r = header_row + 1
    max_row = int(ws.max_row or 1)
    while r <= max_row:
        label = str(ws.cell(r, 1).value or "").strip()
        slot = normalize_time_label(label)
        if slot:
            sessions: dict[str, int] = {}
            corr: dict[str, int] = {}
            for idx, unit in enumerate(units):
                col = idx + 2
                sessions[unit] = _safe_int(ws.cell(r, col).value)
            r += 1
            if r <= max_row and CORR_LABEL_RE.search(str(ws.cell(r, 1).value or "")):
                for idx, unit in enumerate(units):
                    col = idx + 2
                    corr[unit] = _safe_int(ws.cell(r, col).value)
                r += 1
            else:
                corr = {u: 0 for u in units}
            slots[slot] = {"sessions": sessions, "correspondents": corr}
            continue
        r += 1

    if header_row == 1:
        current_slot, sessions, corr = _read_current_period_block(ws, units)
        if current_slot:
            slots[current_slot] = {"sessions": sessions, "correspondents": corr}

    wb.close()
    if not slots:
        raise ValueError("В двухчасовой таблице не найдены временные интервалы")
    return {"units": units, "slots": slots}


def _discover_section(ws, start_row: int, end_row: int) -> dict[str, Any]:
    header_row: int | None = None
    units: list[str] = []
    slots: dict[str, int] = {}
    parts: dict[str, int] = {}
    half_day_rows: list[int] = []
    total_row: int | None = None
    daily_max_row: int | None = None

    for r in range(start_row, end_row + 1):
        label = str(ws.cell(r, 1).value or "").strip()
        if not label:
            continue
        low = label.casefold()
        if low == "бланк":
            header_row = r
            units = _read_header_units(ws, r)
            continue
        slot = normalize_time_label(label)
        if slot:
            slots[slot] = r
            continue
        if low.startswith("часть"):
            parts[label.replace(" ", "").upper()] = r
            continue
        if low.startswith("за 12"):
            half_day_rows.append(r)
            continue
        if low == "всего":
            total_row = r
            continue
        if "максимальное количество корреспондентов" in low:
            daily_max_row = r

    part_groups = _group_parts(slots)
    return {
        "header_row": header_row,
        "units": units,
        "slots": slots,
        "parts": parts,
        "part_groups": part_groups,
        "half_day_rows": half_day_rows,
        "total_row": total_row,
        "daily_max_row": daily_max_row,
    }


def _group_parts(slots: dict[str, int]) -> list[list[str]]:
    ordered = sorted(slots.items(), key=lambda x: x[1])
    labels = [x[0] for x in ordered]
    groups: list[list[str]] = []
    chunk: list[str] = []
    for label in labels:
        chunk.append(label)
        if len(chunk) == 3:
            groups.append(chunk)
            chunk = []
    if chunk:
        groups.append(chunk)
    return groups


def parse_6h_template(content: bytes) -> dict[str, Any]:
    wb = load_workbook(io.BytesIO(content), data_only=False)
    ws = wb.active
    max_row = int(ws.max_row or 1)

    sessions_start = 1
    corr_start = max_row
    for r in range(1, max_row + 1):
        label = str(ws.cell(r, 1).value or "").strip().casefold()
        if "количество корреспондентов" in label:
            corr_start = r
            break

    sessions = _discover_section(ws, sessions_start, corr_start - 1)
    correspondents = _discover_section(ws, corr_start, max_row)

    if not sessions.get("units"):
        raise ValueError("В шестичасовой таблице не найден блок сеансов")
    if not correspondents.get("units"):
        raise ValueError("В шестичасовой таблице не найден блок корреспондентов")

    wb.close()
    return {"sessions": sessions, "correspondents": correspondents, "units": sessions["units"]}


def find_intensity_6h_template_path(launch_root: Path | None = None) -> Path | None:
    import os

    env_path = (os.environ.get("WEB_PORTAL_INTENSITY_6H_TEMPLATE") or "").strip()
    if env_path:
        explicit = Path(env_path).expanduser().resolve()
        if explicit.is_file():
            return explicit

    roots: list[Path] = []
    if launch_root is not None:
        roots.append(Path(launch_root).expanduser().resolve())
    roots.append(Path(BASE_DIR).resolve())
    intensity_dir = Path(BASE_DIR).resolve() / "интенсивность"
    if intensity_dir.is_dir():
        roots.insert(0, intensity_dir)

    seen: set[str] = set()
    candidates: list[Path] = []
    for base in roots:
        key = str(base).lower()
        if key in seen or not base.is_dir():
            continue
        seen.add(key)
        try:
            for item in base.iterdir():
                if not item.is_file() or item.suffix.lower() != ".xlsx":
                    continue
                if item.name.startswith("~$"):
                    continue
                low = item.name.casefold()
                if "6" in low and ("час" in low or "hour" in low or "6_" in low):
                    candidates.append(item.resolve())
        except Exception:
            continue

    if not candidates:
        return None
    candidates.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    return candidates[0]


def _merge_unit_values_sum(existing: dict[str, int], incoming: dict[str, int]) -> dict[str, int]:
    """Сумма по точному названию подразделения (один файл — одна строка)."""
    out = dict(existing)
    for unit, val in incoming.items():
        out[unit] = int(out.get(unit, 0)) + _safe_int(val)
    return out


def _collect_source_data(
    files_2h: list[bytes],
) -> tuple[list[str], dict[str, dict[str, Any]], list[str]]:
    all_units: list[str] = []
    merged_slots: dict[str, dict[str, Any]] = {}
    duplicate_slots: list[str] = []
    unit_seen: set[str] = set()

    for content in files_2h:
        parsed = parse_2h_workbook(content, current_period_only=True)
        for unit in parsed["units"]:
            key = unit.casefold()
            if key not in unit_seen:
                unit_seen.add(key)
                all_units.append(unit)
        for slot, payload in parsed["slots"].items():
            if slot not in merged_slots:
                merged_slots[slot] = {
                    "sessions_files": [dict(payload["sessions"])],
                    "correspondents_files": [dict(payload["correspondents"])],
                    "sessions": dict(payload["sessions"]),
                    "correspondents": dict(payload["correspondents"]),
                }
                continue
            duplicate_slots.append(slot)
            merged_slots[slot]["sessions_files"].append(dict(payload["sessions"]))
            merged_slots[slot]["correspondents_files"].append(dict(payload["correspondents"]))
            merged_slots[slot]["sessions"] = _merge_unit_values_sum(
                merged_slots[slot]["sessions"], payload["sessions"]
            )
            merged_slots[slot]["correspondents"] = _merge_unit_values_sum(
                merged_slots[slot]["correspondents"], payload["correspondents"]
            )

    return all_units, merged_slots, sorted(set(duplicate_slots))


def _value_for_6h_unit(
    slot_entry: dict[str, Any] | None,
    unit_6h: str,
    field: str,
) -> int:
    """
    Сумма по всем загруженным 2-часовым файлам для колонки 6-часовой таблицы.
    В каждом файле берётся лучшее совпадение по названию подразделения.
    """
    if not slot_entry:
        return 0
    files_key = "sessions_files" if field == "sessions" else "correspondents_files"
    file_buckets: list[dict[str, int]] = slot_entry.get(files_key) or []
    total = 0
    for bucket in file_buckets:
        best_val = 0
        best_score = 0.0
        for src, val in bucket.items():
            score = _unit_match_score(unit_6h, src)
            if score >= 40.0 and score > best_score:
                best_score = score
                best_val = _safe_int(val)
        total += best_val
    return total


def _aggregate(values: list[int], mode: str) -> int:
    if not values:
        return 0
    if mode == "sum":
        return sum(values)
    return max(values)


def _write_section_values(
    ws,
    section: dict[str, Any],
    target_units: list[str],
    slot_data: dict[str, dict[str, Any]],
    *,
    field: str,
    aggregate: str,
) -> None:
    units_count = len(section["units"])
    slot_rows = section["slots"]
    updated_slots = set(slot_data.keys())

    for slot, row in slot_rows.items():
        if slot not in updated_slots:
            continue
        entry = slot_data.get(slot)
        for col_idx in range(units_count):
            col = col_idx + 2
            unit_name = target_units[col_idx] if col_idx < len(target_units) else ""
            ws.cell(row, col).value = _value_for_6h_unit(entry, unit_name, field)

    part_rows_sorted = [row for _, row in sorted(section["parts"].items(), key=lambda x: x[1])]
    for gi, group in enumerate(section["part_groups"]):
        if gi >= len(part_rows_sorted):
            break
        row = part_rows_sorted[gi]
        for col_idx in range(units_count):
            col = col_idx + 2
            values = [_safe_int(ws.cell(slot_rows[s], col).value) for s in group if s in slot_rows]
            ws.cell(row, col).value = _aggregate(values, aggregate)

    half_rows = sorted(section.get("half_day_rows") or [])
    if len(half_rows) >= 1 and len(part_rows_sorted) >= 2:
        row = half_rows[0]
        first_half = part_rows_sorted[:2]
        for col_idx in range(units_count):
            col = col_idx + 2
            values = [_safe_int(ws.cell(pr, col).value) for pr in first_half]
            ws.cell(row, col).value = _aggregate(values, aggregate)

    if len(half_rows) >= 2 and len(part_rows_sorted) >= 4:
        row = half_rows[1]
        second_half = part_rows_sorted[2:4]
        for col_idx in range(units_count):
            col = col_idx + 2
            values = [_safe_int(ws.cell(pr, col).value) for pr in second_half]
            ws.cell(row, col).value = _aggregate(values, aggregate)

    total_row = section.get("total_row")
    if total_row and len(half_rows) >= 2:
        for col_idx in range(units_count):
            col = col_idx + 2
            values = [_safe_int(ws.cell(half_rows[i], col).value) for i in range(2)]
            ws.cell(total_row, col).value = _aggregate(values, aggregate)

    daily_max_row = section.get("daily_max_row")
    if daily_max_row and aggregate == "max" and half_rows:
        for col_idx in range(units_count):
            col = col_idx + 2
            values = [_safe_int(ws.cell(r, col).value) for r in half_rows]
            ws.cell(daily_max_row, col).value = _aggregate(values, "max")


def merge_2h_files_into_6h(
    files_2h: list[bytes],
    template_6h: bytes | None = None,
    *,
    launch_root: Path | None = None,
) -> tuple[bytes, dict[str, Any]]:
    if not files_2h:
        raise ValueError("Нужна хотя бы одна двухчасовая таблица")
    if len(files_2h) > 3:
        raise ValueError("Не больше 3 двухчасовых таблиц за раз")

    template_path: Path | None = None
    if template_6h:
        tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        tmp.write(template_6h)
        tmp.close()
        template_path = Path(tmp.name)
    else:
        template_path = find_intensity_6h_template_path(launch_root)

    if template_path is None or not template_path.is_file():
        raise ValueError(
            "Не найден шаблон 6-часовой таблицы. "
            "Положите файл в папку «интенсивность» или укажите WEB_PORTAL_INTENSITY_6H_TEMPLATE."
        )

    source_units, slot_data, duplicate_slots = _collect_source_data(files_2h)
    layout = parse_6h_template(template_path.read_bytes())
    target_units = layout["units"]

    wb = load_workbook(template_path)
    ws = wb.active

    _write_section_values(
        ws,
        layout["sessions"],
        target_units,
        slot_data,
        field="sessions",
        aggregate="sum",
    )
    _write_section_values(
        ws,
        layout["correspondents"],
        target_units,
        slot_data,
        field="correspondents",
        aggregate="sum",
    )

    out_buf = io.BytesIO()
    wb.save(out_buf)
    wb.close()

    units_with_data = 0
    for slot in slot_data:
        for u in target_units:
            if _value_for_6h_unit(slot_data[slot], u, "sessions") > 0 or _value_for_6h_unit(
                slot_data[slot], u, "correspondents"
            ) > 0:
                units_with_data += 1
                break

    meta = {
        "files_count": len(files_2h),
        "slots_updated": sorted(slot_data.keys()),
        "duplicate_slots": duplicate_slots,
        "units_total": len(target_units),
        "units_matched": units_with_data,
        "source_units": len(source_units),
        "merge_mode": "sum_across_files",
    }
    return out_buf.getvalue(), meta

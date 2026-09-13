from __future__ import annotations

import json as _json
import os
import re
import shutil
import tempfile
from copy import copy
from datetime import datetime as _dt
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from web_portal.config import BASE_DIR, DEFAULT_DB_NAME, db_path
from web_portal.lib.seans_db import connect_seans_storage
from web_portal.lib.auth_db import get_position_setting
from web_portal.lib.db import (
    canonical_intensity_unit_key,
    connect,
    ensure_db,
    get_network_intensity_clusters_dual,
    get_network_intensity_clusters_dual_daily,
    get_sessions_favorites,
)
from web_portal.lib.network_analysis_helpers import (
    network_analysis_period_bounds,
    network_favorite_units_from_sessions,
    unit_name_matches,
)
import logging

_log = logging.getLogger("web_portal.lib.network_summary_export")

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_ILLEGAL_XLSX_CHARS = re.compile(r"[\000-\010]|[\013-\014]|[\016-\037]")


def _excel_safe_text(value: Any) -> Any:
    if value is None or isinstance(value, (int, float, bool)):
        return value
    return _ILLEGAL_XLSX_CHARS.sub("", str(value))


_EXPORT_XLSX_NAME_RE = re.compile(
    r"^job_\d+_|^Ucet_intensivnosti_\d{4}-\d{2}-\d{2}_",
    re.IGNORECASE,
)


def _is_intensity_template_candidate(path: Path) -> bool:
    if path.suffix.lower() != ".xlsx":
        return False
    name = path.name
    lower = name.lower()
    if lower.startswith("~$"):
        return False
    if _EXPORT_XLSX_NAME_RE.search(name):
        return False
    return (
        "intensivnosti" in lower
        or "интенсивности" in lower
        or "intensity" in lower
    )


def find_intensity_template_path(launch_root: Path | None = None) -> Path | None:
    """
    Корпоративный шаблон «Учёт интенсивности»:
    - WEB_PORTAL_INTENSITY_TEMPLATE (явный путь)
    - иначе один .xlsx в корне запуска exe / проекта (не tmp, не export_jobs).
    """
    env_path = (os.environ.get("WEB_PORTAL_INTENSITY_TEMPLATE") or "").strip()
    if env_path:
        explicit = Path(env_path).expanduser().resolve()
        if explicit.is_file():
            return explicit

    roots: list[Path] = []
    exe_root = (os.environ.get("WEB_PORTAL_EXE_ROOT") or "").strip()
    if exe_root:
        roots.append(Path(exe_root).expanduser().resolve())
    if launch_root is not None:
        roots.append(Path(launch_root).expanduser().resolve())
    roots.append(Path(BASE_DIR).resolve())

    seen: set[str] = set()
    found: list[Path] = []
    for base in roots:
        key = str(base).lower()
        if key in seen or not base.is_dir():
            continue
        seen.add(key)
        try:
            for item in base.iterdir():
                if not item.is_file():
                    continue
                if _is_intensity_template_candidate(item):
                    found.append(item.resolve())
        except Exception:
            continue

    if not found:
        return None
    found.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    return found[0]


def apply_intensity_template_env(launch_root: Path | None = None) -> Path | None:
    """Подставляет WEB_PORTAL_INTENSITY_TEMPLATE из корня запуска, если не задан вручную."""
    if (os.environ.get("WEB_PORTAL_INTENSITY_TEMPLATE") or "").strip():
        return find_intensity_template_path(launch_root)
    tpl = find_intensity_template_path(launch_root)
    if tpl is not None:
        os.environ["WEB_PORTAL_INTENSITY_TEMPLATE"] = str(tpl)
    return tpl


def _copy_cell_style(src, dst) -> None:
    try:
        dst._style = copy(src._style)
    except Exception:
        dst.font = copy(src.font)
        dst.border = copy(src.border)
        dst.fill = copy(src.fill)
        dst.alignment = copy(src.alignment)
        dst.number_format = src.number_format


def _fill_intensity_template_sheet(
    ws,
    *,
    names: list[str],
    prev_map: dict[str, dict],
    cur_map: dict[str, dict],
    get_map,
    prev_start: _dt.datetime,
    start_clean: _dt.datetime,
    end_clean: _dt.datetime,
    label_fn,
) -> None:
    needed_cols = len(names) + 1
    try:
        for rng in list(ws.merged_cells.ranges):
            if int(rng.max_row) <= 6:
                ws.unmerge_cells(str(rng))
    except Exception:
        _log.debug("_fill_intensity_template_sheet: suppressed error", exc_info=True)

    current = int(ws.max_column or 1)
    if current < needed_cols:
        src_col = max(2, current)
        ws.insert_cols(current + 1, amount=needed_cols - current)
        src_letter = get_column_letter(src_col)
        src_width = ws.column_dimensions[src_letter].width
        for col in range(current + 1, needed_cols + 1):
            col_letter = get_column_letter(col)
            if src_width:
                ws.column_dimensions[col_letter].width = src_width
            for row in range(1, 7):
                _copy_cell_style(ws.cell(row, src_col), ws.cell(row, col))

    max_col = max(needed_cols, int(ws.max_column or needed_cols))
    for row in range(1, 7):
        for col in range(1, max_col + 1):
            ws.cell(row, col).value = None

    ws.cell(1, 1).value = ws.cell(1, 1).value or "Сеть"
    for idx, name in enumerate(names, start=2):
        ws.cell(1, idx).value = _excel_safe_text(name)

    ws.cell(2, 1).value = _excel_safe_text(label_fn(prev_start, start_clean))
    corr_label = ws.cell(3, 1).value or "Кол-во корреспондентов"
    ws.cell(3, 1).value = corr_label
    for idx, name in enumerate(names, start=2):
        ws.cell(2, idx).value = int(get_map(prev_map, name).get("sessions") or 0)
        ws.cell(3, idx).value = int(get_map(prev_map, name).get("correspondents") or 0)

    ws.cell(5, 1).value = _excel_safe_text(label_fn(start_clean, end_clean))
    ws.cell(6, 1).value = corr_label
    for idx, name in enumerate(names, start=2):
        ws.cell(5, idx).value = int(get_map(cur_map, name).get("sessions") or 0)
        ws.cell(6, idx).value = int(get_map(cur_map, name).get("correspondents") or 0)

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["A"].hidden = False
    for col in range(2, needed_cols + 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = 12
        ws.column_dimensions[letter].hidden = False
    for col in range(needed_cols + 1, max_col + 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].hidden = True
        for row in range(1, 7):
            ws.cell(row, col).value = None


def _save_workbook_atomic(wb, out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    tmp_fd, tmp_name = tempfile.mkstemp(suffix=".xlsx", dir=str(out_path.parent))
    os.close(tmp_fd)
    tmp_path = Path(tmp_name)
    try:
        wb.save(str(tmp_path))
        tmp_path.replace(out_path)
    except Exception:
        tmp_path.unlink(missing_ok=True)
        raise


def _parse_order_list(raw) -> list[str]:
    if not raw:
        return []
    if isinstance(raw, list):
        out: list[str] = []
        seen: set[str] = set()
        for x in raw:
            s = str(x or "").strip()
            if not s or s in seen:
                continue
            out.append(s)
            seen.add(s)
        return out
    try:
        parsed = _json.loads(str(raw))
        if isinstance(parsed, list):
            out = []
            seen = set()
            for x in parsed:
                s = str(x or "").strip()
                if not s or s in seen:
                    continue
                out.append(s)
                seen.add(s)
            return out
    except Exception:
        _log.debug("_parse_order_list: suppressed error", exc_info=True)
    return []


def build_network_summary_xlsx(out_path: Path, base_dir: Path, params: dict[str, Any]) -> dict[str, object]:
    position_key = str(params.get("position_name") or "__all__").strip() or "__all__"
    position_filter = None if position_key == "__all__" else position_key
    user_id = int(params.get("user_id") or 0)
    start_s = str(params.get("start") or "").strip()
    end_s = str(params.get("end") or "").strip()
    cluster_order = _parse_order_list(params.get("cluster_order"))
    custom_units_raw = params.get("custom_units") or ""
    favorite_units_raw = params.get("favorite_units") or []
    favorites_only = bool(params.get("favorites_only"))
    if not start_s or not end_s:
        raise ValueError("start/end обязательны")
    try:
        start = _dt.fromisoformat(start_s)
        end = _dt.fromisoformat(end_s)
    except Exception as exc:
        raise ValueError("Неверный формат start/end") from exc
    if end <= start:
        raise ValueError("end должен быть > start")
    bounds = network_analysis_period_bounds(start, end)
    start_clean = bounds["start_clean"]
    end_clean = bounds["end_clean"]
    end_query = bounds["end_query"]
    prev_start = bounds["prev_start"]
    prev_end_query = bounds["prev_end_query"]

    p = db_path(DEFAULT_DB_NAME)
    ensure_db(p)
    main_conn = connect(p)
    seans_conn = connect_seans_storage()
    favorite_units: set[str] = set()
    shared_order: list[str] = []
    cur_clusters: list[dict[str, object]] = []
    prev_clusters: list[dict[str, object]] = []
    try:
        daily_clusters = get_network_intensity_clusters_dual_daily(
            main_conn,
            seans_conn=seans_conn,
            cur_start_dt=start_clean.strftime("%Y-%m-%d %H:%M:%S"),
            cur_end_dt=end_query.strftime("%Y-%m-%d %H:%M:%S"),
            prev_start_dt=prev_start.strftime("%Y-%m-%d %H:%M:%S"),
            prev_end_dt=prev_end_query.strftime("%Y-%m-%d %H:%M:%S"),
            position_name=position_filter,
        )
        if daily_clusters is None:
            cur_clusters, prev_clusters = get_network_intensity_clusters_dual(
                main_conn,
                seans_conn=seans_conn,
                cur_start_dt=start_clean.strftime("%Y-%m-%d %H:%M:%S"),
                cur_end_dt=end_query.strftime("%Y-%m-%d %H:%M:%S"),
                prev_start_dt=prev_start.strftime("%Y-%m-%d %H:%M:%S"),
                prev_end_dt=prev_end_query.strftime("%Y-%m-%d %H:%M:%S"),
                position_name=position_filter,
            )
        else:
            cur_clusters, prev_clusters = daily_clusters

        if favorites_only and user_id:
            try:
                favorites = get_sessions_favorites(main_conn, user_id)
            except Exception:
                favorites = []
            favorite_units = network_favorite_units_from_sessions(main_conn, favorites)
        shared_order = _parse_order_list(
            get_position_setting(
                main_conn,
                position_name=position_key,
                key="analysis_network_shared_order",
            )
        )
    finally:
        try:
            seans_conn.close()
        except Exception:
            _log.debug("build_network_summary_xlsx: suppressed error", exc_info=True)
        main_conn.close()

    if not cluster_order and shared_order:
        cluster_order = shared_order

    if isinstance(favorite_units_raw, list):
        for n in favorite_units_raw:
            s = str(n or "").strip()
            if s:
                favorite_units.add(s)
    elif favorite_units_raw:
        for n in str(favorite_units_raw).split(","):
            s = str(n or "").strip()
            if s:
                favorite_units.add(s)

    def _parse_custom_units(raw) -> list[str]:
        if not raw:
            return []
        if isinstance(raw, list):
            return [str(x or "").strip() for x in raw if str(x or "").strip()]
        try:
            parsed = _json.loads(str(raw))
            if isinstance(parsed, list):
                return [
                    str(x or "").strip() for x in parsed if str(x or "").strip()
                ]
        except Exception:
            _log.debug("_parse_custom_units: suppressed error", exc_info=True)
        return [x.strip() for x in str(raw).split(",") if x.strip()]

    custom_units = _parse_custom_units(custom_units_raw)

    # Как в /api/analysis/network-summary: сначала все колонки из seans, потом фильтр избранного.
    names: list[str] = []
    seen: set[str] = set()
    for lst in (cur_clusters, prev_clusters):
        for c in lst:
            n = str(c.get("unit_name") or "")
            if not n or n in seen:
                continue
            seen.add(n)
            names.append(n)

    lower_map = {str(n).strip().lower(): n for n in names}
    for c in custom_units:
        key = str(c).strip().lower()
        if not key or key in lower_map:
            continue
        names.append(str(c).strip())
        lower_map[key] = str(c).strip()

    if favorites_only and favorite_units:
        fav_keys = {canonical_intensity_unit_key(x) for x in favorite_units}
        names = [
            n
            for n in names
            if canonical_intensity_unit_key(str(n or "")) in fav_keys
        ]

    units_only = _parse_custom_units(params.get("units_only") or params.get("unit_name") or "")
    if units_only:
        names = [
            n
            for n in names
            if any(unit_name_matches(str(n or ""), u) for u in units_only)
        ]

    names = sorted(
        names,
        key=lambda x: str(x or "").strip().lower(),
    )

    parsed_order = _parse_order_list(cluster_order)
    cluster_names_by_lower = {str(n).strip().lower(): str(n).strip() for n in names}
    for lst in (cur_clusters, prev_clusters):
        for c in lst:
            n = str(c.get("unit_name") or "").strip()
            if n:
                cluster_names_by_lower.setdefault(n.lower(), n)

    fav_keys_filter = (
        {canonical_intensity_unit_key(x) for x in favorite_units}
        if favorites_only and favorite_units
        else None
    )

    if parsed_order:
        ordered: list[str] = []
        seen2: set[str] = set()
        for x in parsed_order:
            raw = str(x or "").strip()
            if not raw:
                continue
            canon = cluster_names_by_lower.get(raw.lower(), raw)
            if canon in seen2:
                continue
            ordered.append(canon)
            seen2.add(canon)
        for n in custom_units:
            n = str(n or "").strip()
            if not n or n in seen2:
                continue
            canon = cluster_names_by_lower.get(n.lower(), n)
            if fav_keys_filter is not None and canonical_intensity_unit_key(canon) not in fav_keys_filter:
                continue
            ordered.append(canon)
            seen2.add(canon)
        remaining = sorted(
            [n for n in names if n not in seen2],
            key=lambda x: str(x or "").strip().lower(),
        )
        names = ordered + remaining
    elif not names and parsed_order:
        names = parsed_order[:]

    prev_map = {
        str(x.get("unit_name") or "").strip().lower(): x
        for x in (prev_clusters or [])
        if str(x.get("unit_name") or "").strip()
    }
    cur_map = {
        str(x.get("unit_name") or "").strip().lower(): x
        for x in (cur_clusters or [])
        if str(x.get("unit_name") or "").strip()
    }

    def _get_map(m, name):
        return m.get(str(name).strip().lower(), {})

    # Проверяем, что есть подразделения для экспорта
    if not names:
        has_any_cluster = bool(cur_clusters or prev_clusters)
        if not has_any_cluster:
            raise ValueError(
                "Нет данных сеансов за выбранный период. "
                "Проверьте seans.sqlite и выбранную позицию."
            )
        if favorites_only:
            raise ValueError(
                "Нет избранных подразделений за период. "
                "Добавьте избранное или отключите «Только избранное»."
            )
        if units_only:
            raise ValueError(
                f"Нет данных для подразделения «{units_only[0]}» за выбранный период."
            )
        raise ValueError("Не указан Ucet intensivnosti (нет данных для экспорта)")

    def _label(s: _dt.datetime, e: _dt.datetime) -> str:
        if s.date() == e.date():
            return f"{s:%H.%M}-{e:%H.%M}"
        return f"{s:%d.%m.%Y %H:%M} - {e:%d.%m.%Y %H:%M}"

    template_path = find_intensity_template_path(base_dir)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    if template_path is not None:
        tmp_fd, tmp_name = tempfile.mkstemp(suffix=".xlsx", dir=str(out_path.parent))
        os.close(tmp_fd)
        tmp_path = Path(tmp_name)
        try:
            shutil.copy2(template_path, tmp_path)
            wb = load_workbook(tmp_path)
            ws = wb.active
            _fill_intensity_template_sheet(
                ws,
                names=names,
                prev_map=prev_map,
                cur_map=cur_map,
                get_map=_get_map,
                prev_start=prev_start,
                start_clean=start_clean,
                end_clean=end_clean,
                label_fn=_label,
            )
            _save_workbook_atomic(wb, out_path)
            tmp_path.unlink(missing_ok=True)
            fname = f"Ucet_intensivnosti_{start_clean:%Y-%m-%d_%H-%M}_{end_clean:%Y-%m-%d_%H-%M}.xlsx"
            return {
                "path": str(out_path),
                "filename": fname,
                "mimetype": XLSX_MIME,
                "template": str(template_path),
            }
        except Exception:
            try:
                tmp_path.unlink(missing_ok=True)
            except Exception:
                _log.debug("_label: suppressed error", exc_info=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Ucet Intensivnosti"

    header_fill = PatternFill(
        start_color="FFFF00", end_color="FFFF00", fill_type="solid"
    )
    empty_header_fill = PatternFill(
        start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"
    )
    period_label_fill = PatternFill(
        start_color="FFE4B5", end_color="FFE4B5", fill_type="solid"
    )
    period_data_fill = PatternFill(
        start_color="F2F2F2", end_color="F2F2F2", fill_type="solid"
    )
    correspondents_label_fill = PatternFill(
        start_color="E7E6E6", end_color="E7E6E6", fill_type="solid"
    )
    correspondents_data_fill = PatternFill(
        start_color="F2F2F2", end_color="F2F2F2", fill_type="solid"
    )
    separator_fill = PatternFill(
        start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"
    )

    header_font = Font(bold=True, size=11, name="Calibri")
    data_font = Font(size=11, name="Calibri")
    thin_border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=False)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=False)
    number_format = "0"

    ws.column_dimensions["A"].width = 20
    for col_idx in range(2, len(names) + 2):
        ws.column_dimensions[get_column_letter(col_idx)].width = 12

    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 15
    ws.row_dimensions[3].height = 15
    ws.row_dimensions[4].height = 6
    ws.row_dimensions[5].height = 15
    ws.row_dimensions[6].height = 15

    ws.cell(1, 1).value = "Сеть"
    ws.cell(1, 1).fill = empty_header_fill
    ws.cell(1, 1).font = header_font
    ws.cell(1, 1).border = thin_border
    ws.cell(1, 1).alignment = left_align

    for idx, name in enumerate(names, start=2):
        cell = ws.cell(1, idx)
        cell.value = _excel_safe_text(name)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = center_align

    ws.cell(2, 1).value = _excel_safe_text(_label(prev_start, start_clean))
    ws.cell(2, 1).fill = period_label_fill
    ws.cell(2, 1).font = data_font
    ws.cell(2, 1).border = thin_border
    ws.cell(2, 1).alignment = left_align

    for idx, name in enumerate(names, start=2):
        cell = ws.cell(2, idx)
        cell.value = int(_get_map(prev_map, name).get("sessions") or 0)
        cell.fill = period_data_fill
        cell.font = data_font
        cell.border = thin_border
        cell.alignment = center_align
        cell.number_format = number_format

    ws.cell(3, 1).value = "Кол-во корреспондентов"
    ws.cell(3, 1).fill = correspondents_label_fill
    ws.cell(3, 1).font = data_font
    ws.cell(3, 1).border = thin_border
    ws.cell(3, 1).alignment = left_align

    for idx, name in enumerate(names, start=2):
        cell = ws.cell(3, idx)
        cell.value = int(_get_map(prev_map, name).get("correspondents") or 0)
        cell.fill = correspondents_data_fill
        cell.font = data_font
        cell.border = thin_border
        cell.alignment = center_align
        cell.number_format = number_format

    for col_idx in range(1, len(names) + 2):
        cell = ws.cell(4, col_idx)
        cell.fill = separator_fill
        cell.border = thin_border

    ws.cell(5, 1).value = _excel_safe_text(_label(start_clean, end_clean))
    ws.cell(5, 1).fill = period_label_fill
    ws.cell(5, 1).font = data_font
    ws.cell(5, 1).border = thin_border
    ws.cell(5, 1).alignment = left_align

    for idx, name in enumerate(names, start=2):
        cell = ws.cell(5, idx)
        cell.value = int(_get_map(cur_map, name).get("sessions") or 0)
        cell.fill = period_data_fill
        cell.font = data_font
        cell.border = thin_border
        cell.alignment = center_align
        cell.number_format = number_format

    ws.cell(6, 1).value = "Кол-во корреспондентов"
    ws.cell(6, 1).fill = correspondents_label_fill
    ws.cell(6, 1).font = data_font
    ws.cell(6, 1).border = thin_border
    ws.cell(6, 1).alignment = left_align

    for idx, name in enumerate(names, start=2):
        cell = ws.cell(6, idx)
        cell.value = int(_get_map(cur_map, name).get("correspondents") or 0)
        cell.fill = correspondents_data_fill
        cell.font = data_font
        cell.border = thin_border
        cell.alignment = center_align
        cell.number_format = number_format

    _save_workbook_atomic(wb, out_path)

    fname = f"Ucet_intensivnosti_{start_clean:%Y-%m-%d_%H-%M}_{end_clean:%Y-%m-%d_%H-%M}.xlsx"
    return {"path": str(out_path), "filename": fname, "mimetype": XLSX_MIME}
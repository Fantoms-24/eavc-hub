from __future__ import annotations

from collections import defaultdict
from datetime import datetime, timedelta
from itertools import combinations
import re
import shutil
from pathlib import Path
from typing import Any

import docx
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from web_portal.application.intercepts import build_intercepts_export_docx_file
from web_portal.lib.aviation_word_export import build_aviation_word_docx
from web_portal.lib.callsigns_form_export import build_callsigns_form_xlsx
from web_portal.lib.network_summary_export import build_network_summary_xlsx
from web_portal.config import DEFAULT_DB_NAME, db_path
from web_portal.lib.db import (
    build_sessions_favorite_match_set,
    connect,
    ensure_db,
    get_doc_stats,
    get_seanses_rows_export,
    get_sessions_favorites,
    init_db,
    init_seans_tables,
    sessions_row_is_favorite,
    _build_unit_name_map_for_pairs,
)
from web_portal.lib.document_report import DocumentReport
from web_portal.lib.docx_lang import set_doc_language_ru, set_run_language_ru
from web_portal.lib.export_jobs import export_results_dir


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


from web_portal.lib.network_analysis_helpers import unit_name_matches
import logging

_log = logging.getLogger("web_portal.lib.export_job_builders")


def _pairs_for_unit_filter(main_conn, unit_filter: str) -> set[tuple[str, str]]:
    """Пары (frequency, group_) из справочника unit для фильтра по подразделению."""
    if not unit_filter or main_conn is None:
        return set()
    init_db(main_conn)
    rows = main_conn.execute("SELECT frequency, group_, name FROM unit").fetchall()
    pairs: set[tuple[str, str]] = set()
    for r in rows or []:
        if not unit_name_matches(str(r["name"] or ""), unit_filter):
            continue
        freq = str(r["frequency"] or "").strip()
        grp = str(r["group_"] or "").strip()
        if freq and grp:
            pairs.add((freq, grp))
    return pairs


def _filter_seanses_export_rows(
    rows: list[dict[str, Any]],
    *,
    main_conn,
    unit_filter: str,
) -> list[dict[str, Any]]:
    if not unit_filter:
        return rows
    pairs = _pairs_for_unit_filter(main_conn, unit_filter)
    if not pairs:
        return []
    out: list[dict[str, Any]] = []
    for r in rows:
        freq = str(r.get("frequency") or "").strip()
        grp = str(r.get("group") or r.get("group_") or "").strip()
        if (freq, grp) in pairs:
            out.append(r)
            continue
        # нормализованная частота (151,1750 vs 151.1750)
        try:
            fk = f"{float(freq.replace(',', '.')):.6f}".rstrip("0").rstrip(".")
        except Exception:
            fk = freq
        for pf, pg in pairs:
            try:
                pk = f"{float(pf.replace(',', '.')):.6f}".rstrip("0").rstrip(".")
            except Exception:
                pk = pf
            if pk == fk and pg == grp:
                out.append(r)
                break
    return out


def _default_db_conn():
    p = db_path(DEFAULT_DB_NAME)
    ensure_db(p)
    return connect(p)


def _sessions_export_connections() -> tuple[Any, Any]:
    """main.sqlite — избранное и unit; seans.sqlite — таблицы seanses."""
    from web_portal.lib.seans_db import connect_seans_storage

    main_conn = _default_db_conn()
    seans_conn = connect_seans_storage()
    return main_conn, seans_conn


def _frequency_sort_value(freq: str) -> float:
    try:
        return float(str(freq or "").strip().replace(",", "."))
    except Exception:
        return 0.0


def _group_sort_value(group: str) -> int:
    import re

    m = re.search(r"\d+", str(group or ""))
    if m:
        try:
            return int(m.group(0))
        except Exception:
            return 0
    return 0


def build_sessions_word_export(base_dir: Path, job_id: int, params: dict[str, Any]) -> dict[str, object]:
    datetime_from = str(params.get("datetime_from") or "")
    datetime_to = str(params.get("datetime_to") or "")
    date_from = str(params.get("date_from") or datetime_from[:10] or "from")
    date_to = str(params.get("date_to") or datetime_to[:10] or "to")
    position_key = str(params.get("position_name") or "").strip()
    user_id = int(params.get("user_id") or 0)
    favorites_only = bool(params.get("favorites_only", True))
    unit_filter = str(params.get("unit_name") or params.get("unit_filter") or "").strip()
    out_dir = export_results_dir(base_dir)
    safe_unit = re.sub(r"[^\w\s\-]+", "", unit_filter)[:40].strip().replace(" ", "_") if unit_filter else ""
    filename = (
        f"seanses_{safe_unit}_{date_from[:10]}_{date_to[:10]}.docx"
        if safe_unit
        else f"seanses_{date_from[:10]}_{date_to[:10]}.docx"
    )
    out_path = out_dir / f"job_{job_id}_{filename}"
    main_conn, seans_conn = _sessions_export_connections()
    try:
        if position_key == "__all__":
            table_data = get_doc_stats(
                seans_conn,
                datetime_from,
                datetime_to,
                skip_init=True,
                unit_conn=main_conn,
            )
        else:
            table_data = get_doc_stats(
                seans_conn,
                datetime_from,
                datetime_to,
                client_name=position_key,
                skip_init=True,
                unit_conn=main_conn,
            )
        if favorites_only:
            favorites = get_sessions_favorites(main_conn, user_id) if user_id else []
            fav_keys = build_sessions_favorite_match_set(favorites)
            table_data = [
                row
                for row in table_data
                if sessions_row_is_favorite(
                    str(row.get("frequency") or ""),
                    str(row.get("group") or ""),
                    fav_keys,
                )
            ]
        if unit_filter:
            table_data = [
                row
                for row in table_data
                if unit_name_matches(str(row.get("name") or ""), unit_filter)
            ]
    finally:
        main_conn.close()
        seans_conn.close()
    if not table_data:
        if unit_filter:
            raise ValueError(
                f"Нет данных для подразделения «{unit_filter}» за выбранный период."
            )
        if favorites_only:
            raise ValueError(
                "Нет избранных сетей за выбранный период. Отметьте звёздочкой или снимите «Только избранное»."
            )
        raise ValueError("Нет данных за выбранный период для экспорта")
    report = DocumentReport(datetime_from, datetime_to)
    for row in sorted(
        table_data,
        key=lambda r: (
            -_frequency_sort_value(str(r.get("frequency") or "")),
            -_group_sort_value(str(r.get("group") or "")),
            str(r.get("name") or "").strip().lower(),
        ),
    ):
        report.add_row(
            (
                row.get("frequency", ""),
                row.get("group", ""),
                row.get("name", ""),
                row.get("count", 0),
                row.get("ids_csv", ""),
            )
        )
    report.save(str(out_path))
    return {"path": str(out_path), "filename": filename, "mimetype": DOCX_MIME}


def build_sessions_excel_export(base_dir: Path, job_id: int, params: dict[str, Any]) -> dict[str, object]:
    datetime_from = str(params.get("datetime_from") or "")
    datetime_to = str(params.get("datetime_to") or "")
    date_from = str(params.get("date_from") or datetime_from[:10] or "from")
    date_to = str(params.get("date_to") or datetime_to[:10] or "to")
    position_key = str(params.get("position_name") or "").strip()
    unit_filter = str(params.get("unit_name") or params.get("unit_filter") or "").strip()
    out_dir = export_results_dir(base_dir)
    safe_unit = re.sub(r"[^\w\s\-]+", "", unit_filter)[:40].strip().replace(" ", "_") if unit_filter else ""
    filename = (
        f"seanses_{safe_unit}_{date_from[:10]}_{date_to[:10]}.xlsx"
        if safe_unit
        else f"seanses_{date_from[:10]}_{date_to[:10]}.xlsx"
    )
    out_path = out_dir / f"job_{job_id}_{filename}"
    main_conn, seans_conn = _sessions_export_connections()
    try:
        if position_key == "__all__":
            rows = get_seanses_rows_export(seans_conn, datetime_from, datetime_to)
        else:
            rows = get_seanses_rows_export(
                seans_conn, datetime_from, datetime_to, client_name=position_key
            )
        rows = _filter_seanses_export_rows(rows, main_conn=main_conn, unit_filter=unit_filter)
    finally:
        main_conn.close()
        seans_conn.close()
    if not rows:
        if unit_filter:
            raise ValueError(
                f"Нет строк сеансов для подразделения «{unit_filter}» за выбранный период."
            )
        raise ValueError("Нет данных за выбранный период для экспорта")
    wb = Workbook()
    ws = wb.active
    ws.title = "Сеансы"
    headers = ["Время выхода", "Частота", "ID корреспондента", "Группа", "AES ключа", "Color Voice", "Время выхода (сек)"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for row_idx, r in enumerate(rows, start=2):
        ws.cell(row=row_idx, column=1, value=r.get("date_time") or "")
        ws.cell(row=row_idx, column=2, value=r.get("frequency") or "")
        ws.cell(row=row_idx, column=3, value=r.get("id") or "")
        ws.cell(row=row_idx, column=4, value=r.get("group") or "")
        ws.cell(row=row_idx, column=5, value=r.get("aes_key") or "")
        ws.cell(row=row_idx, column=6, value=r.get("color_voice") or "")
        t = r.get("time_seconds")
        ws.cell(row=row_idx, column=7, value=round(t, 3) if t is not None else "")
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 18
    wb.save(str(out_path))
    return {"path": str(out_path), "filename": filename, "mimetype": XLSX_MIME}


def build_analysis_callsigns_export(base_dir: Path, job_id: int, params: dict[str, Any]) -> dict[str, object]:
    out_dir = export_results_dir(base_dir)
    out_path = out_dir / f"job_{job_id}_formulyar_pozivnyh.xlsx"
    return build_callsigns_form_xlsx(out_path, params)


def build_analysis_network_summary_export(
    base_dir: Path, job_id: int, params: dict[str, Any]
) -> dict[str, object]:
    start_s = str(params.get("start") or "")[:16].replace(":", "-").replace("T", "_")
    end_s = str(params.get("end") or "")[:16].replace(":", "-").replace("T", "_")
    out_dir = export_results_dir(base_dir)
    out_path = out_dir / f"job_{job_id}_Ucet_intensivnosti_{start_s}_{end_s}.xlsx"
    return build_network_summary_xlsx(out_path, base_dir, params)


def _safe_docx_name_part(value: str, fallback: str) -> str:
    cleaned = re.sub(r"[^\w\-\u0400-\u04FF]+", "_", str(value or "").strip())
    cleaned = cleaned.strip("_")[:80]
    return cleaned or fallback


def _month_range_from_ym(ym: str) -> tuple[str, str]:
    month_dt = datetime.strptime(str(ym or "")[:7], "%Y-%m")
    if month_dt.month == 12:
        next_month = datetime(month_dt.year + 1, 1, 1)
    else:
        next_month = datetime(month_dt.year, month_dt.month + 1, 1)
    start_s = month_dt.strftime("%Y-%m-01 00:00:00")
    end_s = (next_month - timedelta(seconds=1)).strftime("%Y-%m-%d %H:%M:%S")
    return start_s, end_s


def _count_seans_in_period(
    seans_conn,
    start_s: str,
    end_s: str,
    client_name: str | None,
) -> int:
    init_seans_tables(seans_conn)
    if client_name:
        row = seans_conn.execute(
            """
            SELECT COUNT(*) FROM (
                SELECT 1 FROM seanses
                WHERE date_time BETWEEN ? AND ? AND client_name = ?
                UNION ALL
                SELECT 1 FROM seanses_archive
                WHERE date_time BETWEEN ? AND ? AND client_name = ?
            )
            """,
            (start_s, end_s, client_name, start_s, end_s, client_name),
        ).fetchone()
    else:
        row = seans_conn.execute(
            """
            SELECT COUNT(*) FROM (
                SELECT 1 FROM seanses
                WHERE date_time BETWEEN ? AND ?
                UNION ALL
                SELECT 1 FROM seanses_archive
                WHERE date_time BETWEEN ? AND ?
            )
            """,
            (start_s, end_s, start_s, end_s),
        ).fetchone()
    return int(row[0] or 0)


def _resolve_full_report_period(
    seans_conn,
    *,
    client_name: str | None,
    report_month: str,
    start_s: str,
    end_s: str,
    position_label: str,
) -> tuple[str, str, str, str, bool]:
    requested = str(report_month or start_s[:7] or "").strip()[:7]
    if requested:
        month_start, month_end = _month_range_from_ym(requested)
        if _count_seans_in_period(seans_conn, month_start, month_end, client_name) > 0:
            return month_start, month_end, requested, requested, False

    client_filter = "AND client_name = ?" if client_name else ""
    params: tuple[Any, ...] = (client_name,) if client_name else ()
    row = seans_conn.execute(
        f"""
        SELECT substr(date_time, 1, 7) AS ym
        FROM (
            SELECT date_time, client_name FROM seanses
            UNION ALL
            SELECT date_time, client_name FROM seanses_archive
        )
        WHERE TRIM(COALESCE(date_time, '')) != '' {client_filter}
        GROUP BY ym
        ORDER BY ym DESC
        LIMIT 1
        """,
        params,
    ).fetchone()
    if not row:
        month_hint = requested or start_s[:7] or "?"
        raise ValueError(
            f"Нет данных для полного отчёта за {month_hint} "
            f"(позиция: {position_label}). В базе нет сеансов за этот период."
        )
    ym = str(row[0] or "")[:7]
    month_start, month_end = _month_range_from_ym(ym)
    return month_start, month_end, ym, requested, ym != requested


def _parse_report_dt(value: object) -> datetime | None:
    raw = str(value or "").strip()
    if not raw:
        return None
    for candidate in (raw, raw.replace("T", " ")):
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
            try:
                return datetime.strptime(candidate[:19], fmt)
            except ValueError:
                continue
    try:
        return datetime.fromisoformat(raw.replace("Z", ""))
    except Exception:
        return None


def _format_duration_seconds(seconds: float | int | None) -> str:
    if seconds is None:
        return "—"
    total = int(round(float(seconds or 0)))
    h, rem = divmod(total, 3600)
    m, s = divmod(rem, 60)
    if h:
        return f"{h} ч {m:02d} мин {s:02d} сек"
    if m:
        return f"{m} мин {s:02d} сек"
    return f"{s} сек"


def _add_doc_paragraph(doc, text: str, *, bold: bool = False, size: int | None = None) -> None:
    p = doc.add_paragraph()
    run = p.add_run(str(text or ""))
    run.bold = bool(bold)
    if size:
        run.font.size = Pt(size)
    set_run_language_ru(run)


def _add_doc_picture(doc, image_path: str | Path, *, width_cm: float = 16.0) -> bool:
    path = Path(image_path)
    if not path.is_file():
        return False
    try:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(path), width=Cm(width_cm))
        set_run_language_ru(run)
        return True
    except Exception:
        return False


def _primary_unit_name(node_units: dict[str, dict[str, int]], cid: str) -> str:
    units = node_units.get(cid) or {}
    if not units:
        return ""
    return max(units.items(), key=lambda x: (int(x[1] or 0), str(x[0])))[0]


def _is_cross_unit_interaction(unit_a: str, unit_b: str) -> bool:
    a = str(unit_a or "").strip()
    b = str(unit_b or "").strip()
    if not a or not b:
        return False
    return a.casefold() != b.casefold()


def _add_doc_table(doc, headers: list[str], rows: list[list[object]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr[i].text = str(title)
        hdr[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in hdr[i].paragraphs[0].runs:
            run.bold = True
            set_run_language_ru(run)
    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data[: len(headers)]):
            cells[i].text = str(value if value is not None else "")
            for p in cells[i].paragraphs:
                for run in p.runs:
                    set_run_language_ru(run)


def _ids_from_csv(value: object) -> set[str]:
    out: set[str] = set()
    for part in str(value or "").split(","):
        token = part.strip()
        if not token:
            continue
        token = token.split(" ", 1)[0].strip()
        if token:
            out.add(token)
    return out


def _row_group(row: dict[str, Any]) -> str:
    return str(row.get("group") or row.get("group_") or "").strip()


def _row_freq(row: dict[str, Any]) -> str:
    return str(row.get("frequency") or "").strip()


def _filter_report_rows(
    rows: list[dict[str, Any]],
    *,
    unit_map: dict[tuple[str, str], str],
    unit_filter: str,
    group_query: str,
    frequency: str,
    group_code: str,
) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for row in rows:
        freq = _row_freq(row)
        group = _row_group(row)
        if frequency and freq != frequency:
            continue
        if group_code and group != group_code:
            continue
        if group_query and group_query.lower() not in group.lower():
            continue
        if unit_filter:
            unit_name = str(unit_map.get((freq, group)) or row.get("name") or "")
            if not unit_name_matches(unit_name, unit_filter):
                continue
        out.append(row)
    return out


def _build_report_graph_metrics(
    rows: list[dict[str, Any]], unit_map: dict[tuple[str, str], str]
) -> dict[str, Any]:
    sessions: dict[tuple[str, str, str], set[str]] = defaultdict(set)
    node_counts: dict[str, int] = defaultdict(int)
    node_units: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    for row in rows:
        dt = str(row.get("date_time") or "").strip()
        freq = _row_freq(row)
        group = _row_group(row)
        cid = str(row.get("id") or "").strip()
        if not dt or not cid:
            continue
        sessions[(dt, freq, group)].add(cid)
    for (_dt, freq, group), ids in sessions.items():
        unit = str(unit_map.get((freq, group)) or "")
        for cid in ids:
            node_counts[cid] += 1
            if unit:
                node_units[cid][unit] += 1

    edge_counts: dict[tuple[str, str], int] = defaultdict(int)
    for ids in sessions.values():
        ids_list = sorted(ids)
        if len(ids_list) < 2:
            continue
        for a, b in combinations(ids_list, 2):
            edge_counts[(a, b)] += 1

    nodes = len(node_counts)
    edges = len(edge_counts)
    adjacency: dict[str, set[str]] = defaultdict(set)
    for a, b in edge_counts:
        adjacency[a].add(b)
        adjacency[b].add(a)
    seen: set[str] = set()
    components = 0
    for cid in node_counts:
        if cid in seen:
            continue
        components += 1
        stack = [cid]
        seen.add(cid)
        while stack:
            cur = stack.pop()
            for nxt in adjacency.get(cur, set()):
                if nxt in seen:
                    continue
                seen.add(nxt)
                stack.append(nxt)
    density = (2 * edges / (nodes * (nodes - 1))) if nodes > 1 else 0.0
    top_hubs = sorted(node_counts.items(), key=lambda x: (-x[1], x[0]))[:15]
    top_bridges = sorted(
        ((cid, len(adjacency.get(cid, set())), node_counts.get(cid, 0)) for cid in node_counts),
        key=lambda x: (-x[1], -x[2], x[0]),
    )[:15]
    top_edges = sorted(edge_counts.items(), key=lambda x: (-x[1], x[0]))[:15]
    interaction_rows = []
    for (a, b), weight in sorted(edge_counts.items(), key=lambda x: (-x[1], x[0])):
        unit_a = _primary_unit_name(node_units, a)
        unit_b = _primary_unit_name(node_units, b)
        if not _is_cross_unit_interaction(unit_a, unit_b):
            continue
        interaction_rows.append(
            {
                "id_a": a,
                "id_b": b,
                "weight": int(weight),
                "unit_a": unit_a,
                "unit_b": unit_b,
            }
        )
    return {
        "nodes": nodes,
        "edges": edges,
        "density": density,
        "components": components,
        "top_hubs": top_hubs,
        "top_bridges": top_bridges,
        "top_edges": top_edges,
        "node_units": node_units,
        "interaction_rows": interaction_rows,
    }


def build_analysis_full_report_docx(
    base_dir: Path, job_id: int, params: dict[str, Any]
) -> dict[str, object]:
    start_s = str(params.get("start") or params.get("datetime_from") or "").replace("T", " ").strip()
    end_s = str(params.get("end") or params.get("datetime_to") or "").replace("T", " ").strip()
    if len(start_s) == 16:
        start_s += ":00"
    if len(end_s) == 16:
        end_s += ":59"
    if not start_s or not end_s:
        raise ValueError("start/end обязательны для полного отчёта")

    position_key = str(params.get("position_name") or "").strip() or "__all__"
    position_filter = None if position_key == "__all__" else position_key
    position_label = "Все позиции" if position_key == "__all__" else position_key
    unit_filter = str(params.get("unit_name") or params.get("unit_query") or "").strip()
    group_query = str(params.get("group_query") or "").strip()
    frequency = str(params.get("frequency") or "").strip()
    group_code = str(params.get("group") or params.get("group_code") or "").strip()
    report_month = str(params.get("report_month") or start_s[:7] or "").strip()

    main_conn, seans_conn = _sessions_export_connections()
    try:
        start_s, end_s, report_month, requested_month, month_auto_resolved = _resolve_full_report_period(
            seans_conn,
            client_name=position_filter,
            report_month=report_month,
            start_s=start_s,
            end_s=end_s,
            position_label=position_label,
        )
        summary_rows = get_doc_stats(
            seans_conn,
            start_s,
            end_s,
            client_name=position_filter,
            skip_init=True,
            unit_conn=main_conn,
        )
        raw_rows = get_seanses_rows_export(
            seans_conn, start_s, end_s, client_name=position_filter
        )
        pairs = {
            (_row_freq(r), _row_group(r))
            for r in raw_rows
            if _row_freq(r) and _row_group(r)
        }
        unit_map = _build_unit_name_map_for_pairs(main_conn, sorted(pairs)) if pairs else {}
        raw_rows = _filter_report_rows(
            raw_rows,
            unit_map=unit_map,
            unit_filter=unit_filter,
            group_query=group_query,
            frequency=frequency,
            group_code=group_code,
        )
        summary_rows = _filter_report_rows(
            summary_rows,
            unit_map=unit_map,
            unit_filter=unit_filter,
            group_query=group_query,
            frequency=frequency,
            group_code=group_code,
        )
    finally:
        main_conn.close()
        seans_conn.close()

    if not raw_rows:
        raise ValueError(
            f"Нет данных для полного отчёта за {report_month} "
            f"(позиция: {position_label})."
        )

    session_count = len(
        {
            (
                str(r.get("date_time") or ""),
                _row_freq(r),
                _row_group(r),
                str(r.get("id") or ""),
            )
            for r in raw_rows
        }
    )
    correspondents = {str(r.get("id") or "").strip() for r in raw_rows if str(r.get("id") or "").strip()}
    frequencies = {_row_freq(r) for r in raw_rows if _row_freq(r)}
    groups = {_row_group(r) for r in raw_rows if _row_group(r)}
    filled_durations = [
        float(r.get("time_seconds") or 0)
        for r in raw_rows
        if r.get("time_seconds") not in (None, "")
    ]
    total_duration = sum(filled_durations) if filled_durations else None
    first_dt = min(str(r.get("date_time") or "") for r in raw_rows if str(r.get("date_time") or ""))
    last_dt = max(str(r.get("date_time") or "") for r in raw_rows if str(r.get("date_time") or ""))

    corr_map: dict[str, dict[str, Any]] = defaultdict(
        lambda: {"sessions": 0, "freqs": set(), "groups": set(), "first": "", "last": "", "duration": 0.0}
    )
    for r in raw_rows:
        cid = str(r.get("id") or "").strip()
        if not cid:
            continue
        c = corr_map[cid]
        c["sessions"] = int(c["sessions"] or 0) + 1
        c["freqs"].add(_row_freq(r))
        c["groups"].add(_row_group(r))
        dt = str(r.get("date_time") or "")
        if dt and (not c["first"] or dt < c["first"]):
            c["first"] = dt
        if dt and (not c["last"] or dt > c["last"]):
            c["last"] = dt
        if r.get("time_seconds") not in (None, ""):
            c["duration"] = float(c["duration"] or 0.0) + float(r.get("time_seconds") or 0)

    graph = _build_report_graph_metrics(raw_rows, unit_map)

    out_dir = export_results_dir(base_dir)
    safe_pos = "all" if position_key == "__all__" else _safe_docx_name_part(position_key, "position")
    safe_month = _safe_docx_name_part(report_month, "month")
    filename = f"full_radio_report_{safe_pos}_{safe_month}.docx"
    out_path = out_dir / f"job_{job_id}_{filename}"

    doc = docx.Document()
    set_doc_language_ru(doc)
    title = doc.add_heading("Полный отчёт по радиосети", level=1)
    for run in title.runs:
        set_run_language_ru(run)
    _add_doc_paragraph(
        doc,
        f"Период: {start_s} — {end_s}. Позиция: {position_label}.",
    )
    if month_auto_resolved and requested_month and requested_month != report_month:
        _add_doc_paragraph(
            doc,
            f"Примечание: за запрошенный месяц {requested_month} данных нет; "
            f"отчёт сформирован за {report_month}.",
        )
    if unit_filter or frequency or group_code or group_query:
        _add_doc_paragraph(
            doc,
            "Фильтры: "
            + "; ".join(
                x
                for x in [
                    f"подразделение/поиск: {unit_filter}" if unit_filter else "",
                    f"частота: {frequency}" if frequency else "",
                    f"группа: {group_code}" if group_code else "",
                    f"поиск группы: {group_query}" if group_query else "",
                ]
                if x
            ),
        )

    doc.add_heading("1. Сводка периода", level=2)
    _add_doc_table(
        doc,
        ["Показатель", "Значение"],
        [
            ["Всего сеансов", session_count],
            ["Уникальных корреспондентов", len(correspondents)],
            ["Частот", len(frequencies)],
            ["Групп", len(groups)],
            ["Первая активность", first_dt],
            ["Последняя активность", last_dt],
            ["Суммарная длительность", _format_duration_seconds(total_duration)],
        ],
    )
    if total_duration is None:
        _add_doc_paragraph(
            doc,
            "Примечание: для импортированных ZIP-метаданных длительность аудио не заполнена; "
            "в отчёте показана активность по времени выхода.",
        )

    doc.add_heading("2. Учёт интенсивности и сеансы", level=2)
    top_summary = sorted(summary_rows, key=lambda r: (-int(r.get("count") or 0), str(r.get("frequency") or "")))[:40]
    _add_doc_table(
        doc,
        ["Частота", "Группа", "Подразделение", "Сеансов", "ID"],
        [
            [
                r.get("frequency", ""),
                r.get("group", ""),
                r.get("name", "") or unit_map.get((str(r.get("frequency") or ""), str(r.get("group") or "")), ""),
                r.get("count", 0),
                r.get("ids_csv", ""),
            ]
            for r in top_summary
        ],
    )

    doc.add_heading("3. Корреспонденты", level=2)
    top_corr = sorted(corr_map.items(), key=lambda x: (-int(x[1]["sessions"] or 0), x[0]))[:50]
    _add_doc_table(
        doc,
        ["ID", "Сеансов", "Частот", "Групп", "Первая активность", "Последняя активность", "Длительность"],
        [
            [
                cid,
                meta["sessions"],
                len(meta["freqs"]),
                len(meta["groups"]),
                meta["first"],
                meta["last"],
                _format_duration_seconds(meta["duration"] if filled_durations else None),
            ]
            for cid, meta in top_corr
        ],
    )

    doc.add_heading("4. Граф радиосети", level=2)
    graph_png_path = str(params.get("graph_png_path") or "").strip()
    if graph_png_path and _add_doc_picture(doc, graph_png_path, width_cm=16.0):
        _add_doc_paragraph(doc, "Снимок графа радиосети с текущими настройками модуля.", bold=False)
    else:
        _add_doc_paragraph(
            doc,
            "Снимок графа недоступен (граф не был открыт или PNG не передан). Ниже — расчётные метрики и связи.",
        )
    _add_doc_table(
        doc,
        ["Метрика", "Значение"],
        [
            ["Узлов", graph["nodes"]],
            ["Связей", graph["edges"]],
            ["Плотность", f"{float(graph['density']):.4f}"],
            ["Компонент", graph["components"]],
        ],
    )
    _add_doc_paragraph(doc, "Топ-хабы:", bold=True)
    _add_doc_table(
        doc,
        ["ID", "Сеансов"],
        [[cid, count] for cid, count in graph["top_hubs"][:15]],
    )
    _add_doc_paragraph(doc, "Топ-мосты:", bold=True)
    _add_doc_table(
        doc,
        ["ID", "Соседей", "Сеансов"],
        [[cid, degree, sessions] for cid, degree, sessions in graph["top_bridges"][:15]],
    )
    _add_doc_paragraph(doc, "Топ-связи:", bold=True)
    _add_doc_table(
        doc,
        ["ID 1", "ID 2", "Вес"],
        [[a, b, weight] for (a, b), weight in graph["top_edges"][:15]],
    )

    doc.add_heading("5. Взаимодействия корреспондентов", level=2)
    _add_doc_paragraph(
        doc,
        "Только взаимодействия между разными подразделениями: совместные выходы в одном сеансе "
        "(одинаковые частота, группа и время), когда ID относятся к разным подразделениям.",
    )
    interaction_rows = list(graph.get("interaction_rows") or [])
    max_interactions = 400
    shown = interaction_rows[:max_interactions]
    _add_doc_table(
        doc,
        ["ID 1", "Подразделение 1", "ID 2", "Подразделение 2", "Совместных сеансов"],
        [
            [
                row["id_a"],
                row.get("unit_a") or "—",
                row["id_b"],
                row.get("unit_b") or "—",
                row["weight"],
            ]
            for row in shown
        ],
    )
    if len(interaction_rows) > max_interactions:
        _add_doc_paragraph(
            doc,
            f"Показаны первые {max_interactions} связей из {len(interaction_rows)}.",
        )
    elif not interaction_rows:
        _add_doc_paragraph(
            doc,
            "Межподразделенческих взаимодействий за период не найдено "
            "(связи внутри одного подразделения в отчёт не включаются).",
        )

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            font = run.font
            if not font.name:
                font.name = "Times New Roman"
            set_run_language_ru(run)
    doc.save(str(out_path))
    png_path = Path(graph_png_path) if graph_png_path else None
    if png_path and png_path.is_file():
        try:
            png_path.unlink()
        except Exception:
            _log.debug("build_analysis_full_report_docx: suppressed error", exc_info=True)
    return {"path": str(out_path), "filename": filename, "mimetype": DOCX_MIME}


def build_aviation_word_export(base_dir: Path, job_id: int, params: dict[str, Any]) -> dict[str, object]:
    frequency_id = int(params.get("frequency_id") or 0)
    work_day = str(params.get("date") or "").strip()
    safe_day = work_day.replace("/", "-") if work_day else "current"
    out_dir = export_results_dir(base_dir)
    out_path = out_dir / f"job_{job_id}_aviation_{frequency_id}_{safe_day}.docx"
    return build_aviation_word_docx(out_path, params)


def build_intercepts_docx_export(base_dir: Path, job_id: int, params: dict[str, Any]) -> dict[str, object]:
    position_name = str(params.get("position_name") or "").strip()
    session_id = int(params.get("session_id") or 0)
    conn = _default_db_conn()
    try:
        tmp_path, filename = build_intercepts_export_docx_file(
            conn, position_name=position_name, session_id=session_id
        )
    finally:
        conn.close()
    out_dir = export_results_dir(base_dir)
    out_path = out_dir / f"job_{job_id}_{filename}"
    shutil.move(str(tmp_path), str(out_path))
    return {"path": str(out_path), "filename": filename, "mimetype": DOCX_MIME}


def build_export_job_result(base_dir: Path, job: dict[str, Any]) -> dict[str, object]:
    job_id = int(job.get("id") or 0)
    job_type = str(job.get("job_type") or "")
    params = job.get("params") if isinstance(job.get("params"), dict) else {}
    builders = {
        "sessions_word": build_sessions_word_export,
        "sessions_excel": build_sessions_excel_export,
        "intercepts_docx": build_intercepts_docx_export,
        "analysis_callsigns": build_analysis_callsigns_export,
        "analysis_full_report": build_analysis_full_report_docx,
        "analysis_network_summary": build_analysis_network_summary_export,
        "aviation_word": build_aviation_word_export,
    }
    builder = builders.get(job_type)
    if builder is None:
        raise ValueError(f"Неизвестный тип export job: {job_type}")
    return builder(base_dir, job_id, params)

from __future__ import annotations

from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from web_portal.config import DEFAULT_DB_NAME, db_path
from web_portal.lib.db import (
    connect,
    count_seanses_by_code,
    ensure_db,
    get_analysis_assignments,
    list_intercept_callsigns,
    list_intercept_catalog,
)


def _msk_now_naive() -> datetime:
    return datetime.now(timezone(timedelta(hours=3))).replace(tzinfo=None)


def _fmt_dt(d: datetime) -> str:
    return d.strftime("%Y-%m-%d %H:%M:%S")


def build_callsigns_form_xlsx(out_path: Path, params: dict[str, Any]) -> dict[str, object]:
    pos = str(params.get("position_name") or "").strip()
    if not pos:
        raise ValueError("position_name обязателен")
    days = int(str(params.get("days") or "7").strip() or "7")
    if days not in {1, 3, 7, 30}:
        days = 7

    end = _msk_now_naive()
    start = end - timedelta(days=days)
    start_s = _fmt_dt(start)
    end_s = _fmt_dt(end)

    p = db_path(DEFAULT_DB_NAME)
    ensure_db(p)
    conn = connect(p)
    try:
        catalog = list_intercept_catalog(conn, pos) or []
        callsigns = list_intercept_callsigns(conn, pos) or []
        cs_map: dict[tuple[str, str, str], dict[str, object]] = {}
        for cs in callsigns:
            key = (
                str(cs.get("frequency") or "").strip(),
                str(cs.get("group_code") or "").strip(),
                str(cs.get("code") or "").strip(),
            )
            cs_map[key] = cs

        units: dict[str, list[dict[str, object]]] = {}
        for row in catalog:
            unit_name = str(row.get("unit_name") or "").strip()
            if not unit_name:
                continue
            units.setdefault(unit_name, []).append(row)
    finally:
        conn.close()

    if not units:
        raise ValueError("Нет данных intercept_catalog для экспорта")

    wb = Workbook()
    ws = wb.active
    ws.title = "Формуляр позывных"

    header_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=16)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="top", wrap_text=True)
    thin = Side(border_style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill(fill_type="solid", fgColor="2563EB")
    unit_fill = PatternFill(fill_type="solid", fgColor="E5E7EB")
    stripe_fill = PatternFill(fill_type="solid", fgColor="F9FAFB")

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
    cell = ws.cell(row=1, column=1, value="Формуляр позывных")
    cell.font = title_font
    cell.alignment = center

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=9)
    ws.cell(row=2, column=1, value=f"Позиция: {pos}").alignment = left

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=9)
    ws.cell(row=3, column=1, value=f"Период: {start_s} — {end_s} МСК").alignment = left

    columns = [
        "Подразделение",
        "Частота",
        "Группа",
        "ID",
        "Позывной",
        "Тег",
        "Опер.д.",
        "Сеансов",
        "Описание",
    ]
    header_row = 5
    max_width: list[int] = [len(col) for col in columns]
    for idx, name in enumerate(columns, start=1):
        c = ws.cell(row=header_row, column=idx, value=name)
        c.font = header_font
        c.alignment = center
        c.border = border
        c.fill = header_fill

    current_row = header_row + 1

    conn = connect(p)
    try:
        for unit_name in sorted(units.keys(), key=lambda s: s.lower()):
            ws.merge_cells(
                start_row=current_row,
                start_column=1,
                end_row=current_row,
                end_column=len(columns),
            )
            uh = ws.cell(row=current_row, column=1, value=unit_name)
            uh.font = Font(bold=True)
            uh.alignment = left
            uh.fill = unit_fill
            uh.border = border
            max_width[0] = max(max_width[0], len(unit_name))
            current_row += 1

            rows_for_unit = units[unit_name]
            pairs: dict[tuple[str, str], list[dict[str, object]]] = {}
            for r in rows_for_unit:
                freq = str(r.get("frequency") or "").strip()
                grp = str(r.get("group_code") or "").strip()
                if not freq or not grp:
                    continue
                pairs.setdefault((freq, grp), []).append(r)

            def _freq_sort_key(pair: tuple[str, str]) -> float:
                try:
                    return float(str(pair[0]).replace(",", "."))
                except Exception:
                    return 0.0

            for (freq, grp), _ in sorted(
                pairs.items(),
                key=lambda x: (_freq_sort_key(x[0]), x[0][1]),
            ):
                grp_disp = str(grp or "")
                if grp_disp and not grp_disp.upper().startswith("G"):
                    grp_disp = f"G{grp_disp}"
                pair_title = f"{freq} {grp_disp}".strip()

                ws.merge_cells(
                    start_row=current_row,
                    start_column=2,
                    end_row=current_row,
                    end_column=len(columns),
                )
                ph = ws.cell(row=current_row, column=2, value=pair_title)
                ph.font = Font(bold=True)
                ph.alignment = left
                ph.fill = PatternFill(fill_type="solid", fgColor="DBEAFE")
                for col_idx in range(2, len(columns) + 1):
                    ws.cell(row=current_row, column=col_idx).border = border
                current_row += 1

                counts = count_seanses_by_code(
                    conn,
                    position_name=pos,
                    frequency=freq,
                    group_code=grp,
                    start_dt=start_s,
                    end_dt=end_s,
                )
                assignments = get_analysis_assignments(
                    conn, position_name=pos, frequency=freq, group_code=grp
                )
                duty_value = assignments.get("duty") or []
                if isinstance(duty_value, list):
                    duty_codes = {str(x).strip() for x in duty_value if str(x).strip()}
                else:
                    duty_codes = (
                        {str(duty_value).strip()} if str(duty_value).strip() else set()
                    )

                codes: set[str] = set(str(c) for c in counts.keys())
                for key, cs in cs_map.items():
                    cf, cg, cc = key
                    if cf == freq and cg == grp and cc:
                        codes.add(cc)

                if not codes:
                    continue

                def _sort_key(code: str) -> tuple[int, str]:
                    cnt = int(counts.get(code, 0) or 0)
                    return (-cnt, code)

                for code in sorted(codes, key=_sort_key):
                    cnt = int(counts.get(code, 0) or 0)
                    cs = cs_map.get((freq, grp, code), {}) or {}
                    label = str(cs.get("label") or "")
                    tag = str(cs.get("tag") or "")
                    tag_desc = str(cs.get("tag_desc") or "")
                    tag_color = str(cs.get("tag_color") or "")
                    is_duty = bool(code and code in duty_codes)

                    values = [
                        "",
                        freq,
                        grp,
                        code,
                        label or "—",
                        tag or "—",
                        "Да" if is_duty else "—",
                        cnt,
                        tag_desc or "—",
                    ]
                    for col_idx, val in enumerate(values, start=1):
                        c = ws.cell(row=current_row, column=col_idx, value=val)
                        c.border = border
                        if col_idx in {1, 5, 9}:
                            c.alignment = left
                        elif col_idx in {3, 4, 8}:
                            c.alignment = Alignment(horizontal="right", vertical="center")
                        else:
                            c.alignment = center
                        text_val = str(val)
                        max_width[col_idx - 1] = max(max_width[col_idx - 1], len(text_val))

                    tag_color_valid = (
                        isinstance(tag_color, str)
                        and tag_color.startswith("#")
                        and len(tag_color) == 7
                    )
                    if tag_color_valid:
                        rgb = tag_color.lstrip("#")
                        tag_fill = PatternFill(fill_type="solid", fgColor=rgb)
                        tag_cell = ws.cell(row=current_row, column=6)
                        tag_cell.fill = tag_fill
                        tag_cell.font = Font(bold=True, color="FFFFFF")

                    if (current_row - header_row) % 2 == 0:
                        for col_idx in range(1, len(columns) + 1):
                            if col_idx == 6 and tag_color_valid:
                                continue
                            ws.cell(row=current_row, column=col_idx).fill = stripe_fill

                    current_row += 1
    finally:
        conn.close()

    for idx, width in enumerate(max_width, start=1):
        col_letter = get_column_letter(idx)
        ws.column_dimensions[col_letter].width = min(max(width + 2, 10), 60)

    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:I{max(current_row - 1, header_row)}"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))
    return {
        "path": str(out_path),
        "filename": "formulyar_pozivnyh.xlsx",
        "mimetype": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }

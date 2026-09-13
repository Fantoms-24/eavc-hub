from __future__ import annotations

import docx
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from web_portal.lib.docx_lang import set_doc_language_ru, set_run_language_ru


class DocumentReport:
    def __init__(self, time_start: str, time_end: str) -> None:
        # Извлекаем только время (HH:MM) из datetime строки
        self.time_start = self._extract_time(time_start)
        self.time_end = self._extract_time(time_end)
        self.rows_counter = 0
        self.doc = docx.Document()
        # Русский язык по умолчанию (для орфографии/Т9 в OnlyOffice)
        set_doc_language_ru(self.doc)

        header = [
            "",
            "Принадлежность",
            "Содержание перехвата",
            "Примечания",
        ]
        self.table = self.doc.add_table(rows=1, cols=len(header))
        row = self.table.rows[0].cells
        for i, e in enumerate(header):
            row[i].text = e
            row[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT

    @staticmethod
    def _extract_time(datetime_str: str) -> str:
        """Извлекает только время (HH:MM) из строки datetime для Word."""
        if not datetime_str:
            return ""
        # Формат может быть: "2024-01-15 14:00:00" или "2024-01-15T14:00:00"
        # Преобразуем в формат "14:00"
        import re
        from datetime import datetime

        try:
            # Пробуем разные форматы
            for fmt in ["%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M"]:
                try:
                    dt = datetime.strptime(datetime_str, fmt)
                    return dt.strftime("%H:%M")
                except ValueError:
                    continue
            # Если не удалось распарсить, пытаемся извлечь время
            match = re.search(r"(\d{1,2}):(\d{2})", datetime_str)
            if match:
                return match.group(0)  # Возвращает HH:MM
            return datetime_str  # Если не найдено, возвращаем как есть
        except Exception:
            return datetime_str

    @staticmethod
    def plural(n) -> str:
        try:
            n = int(n)
        except (ValueError, TypeError):
            return f"отмечено\n{n} сеансов"

        if n > 9 and 10 <= n % 100 <= 20:
            return f"отмечено\n{n} сеансов"

        if n % 10 == 1:
            return f"отмечен\n{n} сеанс"
        if n % 10 in {2, 3, 4}:
            return f"отмечено\n{n} сеанса"
        if n % 10 in {5, 6, 7, 8, 9}:
            return f"отмечено\n{n} сеансов"
        return f"отмечено\n{n} сеансов"

    def change_font(self):
        for row in self.table.rows:
            for cell in row.cells:
                if not cell.paragraphs or not cell.paragraphs[0].runs:
                    continue
                run = cell.paragraphs[0].runs[0]
                font = run.font
                font.name = "Times New Roman"
                font.size = Pt(12)
                set_run_language_ru(run)

    def change_width(self):
        self.table.autofit = False
        self.table.allow_autofit = False  # type: ignore
        widths = (
            Cm(1),
            Cm(2.5),
            Cm(8),
            Cm(3),
        )
        for col_idx, width in enumerate(widths):
            self.table.columns[col_idx].width = width
        for row in self.table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = width

    def add_row(self, row_data: tuple):
        """
        Ожидается формат как в desktop:
        (frequency, group_, unit_name, count, ids_csv)
        """
        row = self.table.add_row().cells
        self.rows_counter += 1
        row[0].text = f"{self.rows_counter}."

        group_val = row_data[1] if len(row_data) > 1 else ""
        # Получаем ID корреспондентов: в desktop отрезают ведущую "G"
        correspondent_ids_str = (
            group_val[1:] if isinstance(group_val, str) and len(group_val) > 1 else ""
        )

        notes_str = row_data[4] if len(row_data) > 4 else ""
        if notes_str:
            notes_list = [
                note.strip() for note in str(notes_str).split(",") if note.strip()
            ]
            notes_count = len(notes_list)
        else:
            notes_count = 0

        unit_name = row_data[2] if len(row_data) > 2 else ""
        freq = row_data[0] if len(row_data) > 0 else ""
        count = row_data[3] if len(row_data) > 3 else ""

        row[1].text = f"{unit_name}\n{freq}\nID\n{correspondent_ids_str}"
        row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        row[2].text = (
            f"В период с {self.time_start} до {self.time_end} в р/с {DocumentReport.plural(count)} связи"
        )
        row[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        notes_text = "\n".join(str(notes_str).split(",")) if notes_str else ""
        if notes_text:
            row[3].text = f"{notes_text}\n\nИтого: {notes_count}"
        else:
            row[3].text = f"Итого: {notes_count}"
        row[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    def save(self, filename: str):
        self.table.style = "Table Grid"
        set_doc_language_ru(self.doc)
        self.change_font()
        self.change_width()
        self.doc.save(filename)

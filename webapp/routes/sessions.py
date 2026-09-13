"""Маршруты сеансов (/sessions, /api/sessions/*), извлечены из app.py дословно."""

from __future__ import annotations

import json
import logging
import os
import re
import tempfile
import time
from io import BytesIO
from pathlib import Path

from flask import jsonify, render_template, request, send_file
from flask_login import current_user, login_required
from openpyxl import Workbook

from web_portal.app_perf import perf_route as _perf_route
from web_portal.app_runtime_caches import (
    folder_import_progress_get as _folder_import_progress_get,
    folder_import_progress_set as _folder_import_progress_set,
)
from web_portal.config import (
    DEFAULT_DB_NAME,
    db_path,
    seans_watch_status_path,
    search_online_db_path,
)
from web_portal.application.sessions import (
    SessionsUseCaseHTTP,
    assemble_sessions_favorites_list,
    build_sessions_export_datetime_window,
    execute_folder_import,
    execute_sessions_add_favorite,
    execute_sessions_remove_favorite,
)
from web_portal.lib.db import (
    _normalize_frequency_key,
    build_sessions_favorite_match_set,
    connect,
    ensure_db,
    get_db_stats,
    get_doc_stats,
    get_seanses_rows_export,
    get_sessions_favorites,
    sessions_row_is_favorite,
    update_unit_name,
)
from web_portal.lib.document_report import DocumentReport
from web_portal.lib.flask_db import get_request_main_db, get_request_seans_db
from web_portal.lib.heavy_workers import (
    submit_folder_import as _submit_folder_import,
)
from web_portal.lib.portal_event_log import log_portal_event as _log_portal_event
from web_portal.positions import (
    can_edit_session_unit_name,
    get_selected_position,
    require_position_role,
)
from web_portal.webapp.auth import require_perm, require_tab
from web_portal.webapp.context import AppContext
from web_portal.webapp.httputils import json_etag_response as _json_etag_response
from web_portal.webapp.seans_watch import (
    _clear_seans_watch_autostart,
    _launch_global_seans_watch_worker,
)

_log = logging.getLogger("web_portal.webapp.routes.sessions")


def register_sessions_routes(app, ctx: AppContext):
    """Регистрирует маршруты сеансов через AppContext. Код перенесён из create_app дословно."""

    @app.get("/sessions")
    @login_required
    @require_tab("tab_sessions")
    def sessions():
        return render_template("sessions.html")

    @app.post("/api/sessions/process-folder")
    @require_perm("import_folder")
    def api_sessions_process_folder():
        """Ставит импорт папки в фоновый воркер (не блокирует поток Waitress)."""
        position_name = get_selected_position()
        if not position_name:
            return (
                jsonify(
                    {
                        "ok": False,
                        "error": "Позиция не выбрана (выберите позицию слева)",
                    }
                ),
                400,
            )
        data = request.get_json(silent=True) or {}
        folder_path = (data.get("folder_path") or "").strip()

        if not folder_path:
            return jsonify({"ok": False, "error": "Путь к папке не указан"}), 400

        folder = Path(folder_path)
        if not folder.exists() or not folder.is_dir():
            return jsonify({"ok": False, "error": "Папка не существует"}), 400

        import_id = str(data.get("import_id") or "").strip()
        if not import_id:
            import_id = f"seans-{int(time.time() * 1000)}-{int(current_user.id)}"

        _folder_import_progress_set(
            import_id,
            {
                "user_id": int(current_user.id),
                "phase": "queued",
                "message": "В очереди на импорт…",
                "done": False,
            },
        )

        accepted = _submit_folder_import(
            lambda: execute_folder_import(
                import_id=import_id,
                folder_path=folder_path,
                position_name=position_name,
                user_id=int(current_user.id),
                progress_set=_folder_import_progress_set,
                log_event=_log_portal_event,
            )
        )
        if not accepted:
            _folder_import_progress_set(
                import_id,
                {
                    "phase": "error",
                    "message": "Импорт уже выполняется. Дождитесь завершения.",
                    "done": True,
                    "error": "import_busy",
                },
            )
            return (
                jsonify(
                    {
                        "ok": False,
                        "error": "Импорт папки уже выполняется. Повторите позже.",
                        "import_id": import_id,
                    }
                ),
                503,
            )

        return (
            jsonify(
                {
                    "ok": True,
                    "queued": True,
                    "import_id": import_id,
                    "message": "Импорт запущен в фоне",
                }
            ),
            202,
        )

    @app.get("/api/sessions/process-folder/progress")
    @require_perm("import_folder")
    def api_sessions_process_folder_progress():
        import_id = str(request.args.get("import_id") or "").strip()
        if not import_id:
            return jsonify({"ok": False, "error": "import_id обязателен"}), 400
        progress = _folder_import_progress_get(import_id)
        if not progress:
            return jsonify(
                {
                    "ok": True,
                    "progress": {"phase": "idle", "done": False},
                }
            )
        owner_id = int(progress.get("user_id") or 0)
        if owner_id and owner_id != int(current_user.id):
            return jsonify({"ok": False, "error": "Нет доступа к прогрессу импорта"}), 403
        progress = dict(progress)
        progress.pop("user_id", None)
        progress.pop("updated_ts", None)
        if progress.get("done"):
            return jsonify(
                {
                    "ok": True,
                    "progress": progress,
                    "result": {
                        "total_files": progress.get("total_files", 0),
                        "total_records": progress.get("total_records", 0),
                        "errors": progress.get("errors"),
                        "position_name": progress.get("position_name"),
                    },
                }
            )
        return jsonify({"ok": True, "progress": progress})

    @app.get("/api/sessions/stats")
    @_perf_route("api/sessions/stats")
    @login_required
    def api_sessions_stats():
        """Получает статистику по сеансам в БД"""
        try:
            conn = get_request_seans_db()
            stats = get_db_stats(conn)
            return jsonify({"ok": True, "stats": stats, "cache_hit": False})
        except Exception as e:
            import logging

            logging.getLogger("web_portal.sessions").error(f"Error getting stats: {e}")
            return jsonify({"ok": False, "error": str(e)}), 500

    @app.get("/api/sessions/table-data")
    @_perf_route("api/sessions/table-data")
    @login_required
    def api_sessions_table_data():
        """Получает данные для таблицы сеансов за указанный период"""
        try:
            import time as _time

            window = build_sessions_export_datetime_window(
                date_from=request.args.get("date_from", "").strip(),
                time_from=request.args.get("time_from", "00:00").strip(),
                date_to=request.args.get("date_to", "").strip(),
                time_to=request.args.get("time_to", "23:59").strip(),
            )
            datetime_from = window.datetime_from
            datetime_to = window.datetime_to

            seans_conn = get_request_seans_db()
            main_conn = get_request_main_db()
            t_handler = _time.perf_counter()
            position_name = get_selected_position()
            # Админ: если позиция не выбрана — показываем общие данные со всех позиций
            if current_user.role == "admin" and not position_name:
                all_positions = True
            else:
                if not position_name:
                    return (
                        jsonify(
                            {
                                "ok": False,
                                "error": "Позиция не выбрана (выберите позицию слева)",
                            }
                        ),
                        400,
                    )
                all_positions = False

            t_sql = _time.perf_counter()
            page_data = get_doc_stats(
                seans_conn,
                datetime_from,
                datetime_to,
                client_name=None if all_positions else position_name,
                unit_conn=main_conn,
            )
            sql_ms = int((_time.perf_counter() - t_sql) * 1000)
            total = len(page_data)

            # Получаем избранное пользователя (не кэшируется — per-user)
            favorites = []
            if current_user.is_authenticated:
                favorites = get_sessions_favorites(main_conn, int(current_user.id))

            # Помечаем избранные записи (та же группа + зона частоты, не точная строка)
            fav_keys = build_sessions_favorite_match_set(favorites)
            for item in page_data:
                item["is_favorite"] = sessions_row_is_favorite(
                    str(item.get("frequency") or ""),
                    str(item.get("group") or ""),
                    fav_keys,
                )

            total_ms = int((_time.perf_counter() - t_handler) * 1000)

            from web_portal.lib.seans_db import describe_seans_storage_resolution

            storage_info = describe_seans_storage_resolution()

            return _json_etag_response(
                {
                    "ok": True,
                    "data": page_data,
                    "total": total,
                    "period_from": datetime_from,
                    "period_to": datetime_to,
                    "favorites": favorites,
                    "sql_ms": sql_ms,
                    "total_ms": total_ms,
                    "query_ms": total_ms,
                    "cache_hit": False,
                    "position_name": position_name,
                    "all_positions": all_positions,
                    "storage_path": storage_info.get("active_path"),
                    "storage_rows": {
                        "main": storage_info.get("main_rows"),
                        "seans": storage_info.get("seans_rows"),
                    },
                },
                max_age=0,
            )
        except Exception as e:
            import logging

            logging.getLogger("web_portal.sessions").error(
                f"Error getting table data: {e}"
            )
            return jsonify({"ok": False, "error": str(e)}), 500

    @app.post("/api/sessions/start-watch")
    @require_perm("import_folder")
    def api_sessions_start_watch():
        """Запускает автопоиск новых файлов result0.txt"""
        position_name = get_selected_position()
        if not position_name:
            return (
                jsonify(
                    {
                        "ok": False,
                        "error": "Позиция не выбрана (выберите позицию слева)",
                    }
                ),
                400,
            )
        data = request.get_json(silent=True) or {}
        folder_path = (data.get("folder_path") or "").strip()
        interval_sec = float(data.get("interval_sec") or 5)
        if interval_sec < 1:
            interval_sec = 1
        if interval_sec > 60:
            interval_sec = 60

        if not folder_path:
            return jsonify({"ok": False, "error": "Путь к папке не указан"}), 400

        folder = Path(folder_path)
        if not folder.exists() or not folder.is_dir():
            return jsonify({"ok": False, "error": "Папка не существует"}), 400

        try:
            out = _launch_global_seans_watch_worker(
                app, position_name, folder_path, interval_sec
            )
            return jsonify(out)
        except Exception as e:
            import logging

            logging.getLogger("web_portal.sessions").error(f"Error starting watch: {e}")
            return jsonify({"ok": False, "error": str(e)}), 500

    @app.post("/api/sessions/stop-watch")
    @login_required
    def api_sessions_stop_watch():
        """Останавливает автопоиск"""
        try:
            if not hasattr(app, "session_watchers"):
                return jsonify({"ok": True, "message": "Автопоиск не запущен"})

            watcher_key = "hub_global"

            if watcher_key in app.session_watchers:
                w = app.session_watchers[watcher_key]
                if w.get("mode") == "subprocess":
                    sp = w.get("stop_path")
                    proc = w.get("proc")
                    if sp:
                        try:
                            sp.touch()
                        except Exception:
                            _log.debug("api_sessions_stop_watch: suppressed error", exc_info=True)
                    if proc:
                        try:
                            proc.terminate()
                            proc.wait(timeout=8)
                        except Exception:
                            try:
                                proc.kill()
                            except Exception:
                                _log.debug("api_sessions_stop_watch: suppressed error", exc_info=True)
                    try:
                        seans_watch_status_path().unlink(missing_ok=True)
                    except Exception:
                        _log.debug("api_sessions_stop_watch: suppressed error", exc_info=True)
                else:
                    try:
                        w["stop"].set()
                    except Exception:
                        _log.debug("api_sessions_stop_watch: suppressed error", exc_info=True)
                    try:
                        thr = w.get("thread")
                        if thr:
                            thr.join(timeout=5)
                    except Exception:
                        _log.debug("api_sessions_stop_watch: suppressed error", exc_info=True)
                del app.session_watchers[watcher_key]
                _clear_seans_watch_autostart()
                return jsonify({"ok": True, "message": "Автопоиск остановлен"})
            _clear_seans_watch_autostart()
            return jsonify({"ok": True, "message": "Автопоиск не запущен"})
        except Exception as e:
            import logging

            logging.getLogger("web_portal.sessions").error(f"Error stopping watch: {e}")
            return jsonify({"ok": False, "error": str(e)}), 500

    @app.get("/api/sessions/watch-status")
    @login_required
    def api_sessions_watch_status():
        """Статус автопоиска (последняя проверка / обработано / ошибка)"""
        try:
            if not hasattr(app, "session_watchers"):
                return jsonify({"ok": True, "status": {"running": False}})

            watcher_key = "hub_global"
            w = app.session_watchers.get(watcher_key)
            if not w:
                return jsonify({"ok": True, "status": {"running": False}})

            if w.get("mode") == "subprocess":
                p = w.get("status_path") or seans_watch_status_path()
                proc = w.get("proc")
                if p.exists():
                    try:
                        st = json.loads(p.read_text(encoding="utf-8"))
                        if proc is not None and proc.poll() is not None:
                            st["running"] = False
                        return jsonify({"ok": True, "status": st})
                    except Exception:
                        _log.debug("api_sessions_watch_status: suppressed error", exc_info=True)
                run = proc is not None and proc.poll() is None
                return jsonify(
                    {
                        "ok": True,
                        "status": {"running": run, "mode": "subprocess", "last_message": "нет файла статуса"},
                    }
                )

            status = w.get("status") or {}
            lock = w.get("status_lock")
            if lock:
                try:
                    with lock:
                        return jsonify({"ok": True, "status": dict(status)})
                except Exception:
                    _log.debug("api_sessions_watch_status: suppressed error", exc_info=True)
            return jsonify({"ok": True, "status": dict(status)})
        except Exception as e:
            import logging

            logging.getLogger("web_portal.sessions").error(
                f"Error getting watch status: {e}"
            )
            return jsonify({"ok": False, "error": str(e)}), 500

    @app.post("/api/sessions/export-word")
    @_perf_route("api/sessions/export-word")
    @login_required
    def api_sessions_export_word():
        """Экспортирует данные сеансов в Word документ"""
        try:
            from flask import after_this_request

            data = request.get_json(silent=True) or {}
            window = build_sessions_export_datetime_window(
                date_from=data.get("date_from", "").strip(),
                time_from=data.get("time_from", "00:00").strip(),
                date_to=data.get("date_to", "").strip(),
                time_to=data.get("time_to", "23:59").strip(),
            )
            datetime_from = window.datetime_from
            datetime_to = window.datetime_to
            date_from = window.date_from
            date_to = window.date_to

            p = db_path(DEFAULT_DB_NAME)
            ensure_db(p)
            main_conn = get_request_main_db()
            seans_conn = get_request_seans_db()
            position_name = get_selected_position()
            if current_user.role == "admin" and not position_name:
                table_data = get_doc_stats(
                    seans_conn,
                    datetime_from,
                    datetime_to,
                    unit_conn=main_conn,
                )
            else:
                if not position_name:
                    return (
                        jsonify(
                            {
                                "ok": False,
                                "error": "Позиция не выбрана (выберите позицию слева)",
                            }
                        ),
                        400,
                    )
                if not require_position_role(position_name, "export"):
                    return jsonify({"ok": False, "error": "Нет прав на экспорт"}), 403
                table_data = get_doc_stats(
                    seans_conn,
                    datetime_from,
                    datetime_to,
                    client_name=position_name,
                    unit_conn=main_conn,
                )

            favorites = get_sessions_favorites(main_conn, int(current_user.id))
            fav_keys = build_sessions_favorite_match_set(favorites)
            table_data = [
                row
                for row in table_data
                if sessions_row_is_favorite(
                    str(row.get("frequency") or ""),
                    str(row.get("group") or ""),
                    fav_keys,
                )
            ]
            table_data = sorted(
                table_data,
                key=lambda r: (
                    str(r.get("name") or "").strip().lower(),
                    str(r.get("frequency") or "").strip(),
                    str(r.get("group") or "").strip(),
                ),
            )

            if not table_data:
                return (
                    jsonify(
                        {"ok": False, "error": "Нет данных избранного для экспорта"}
                    ),
                    400,
                )

            report = DocumentReport(datetime_from, datetime_to)

            # Добавляем данные в отчет
            for row in table_data:
                # row это dict с ключами: frequency, group, name, count, ids_csv
                freq = row.get("frequency", "")
                group = row.get("group", "")
                unit_name = row.get("name", "")
                count = row.get("count", 0)
                ids = row.get("ids_csv", "")

                report.add_row((freq, group, unit_name, count, ids))

            # Сохраняем во временный файл
            temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
            temp_path = temp_file.name
            temp_file.close()

            report.save(temp_path)

            @after_this_request
            def _cleanup_session_word(response):
                try:
                    os.unlink(temp_path)
                except Exception:
                    _log.debug("_cleanup_session_word: suppressed error", exc_info=True)
                return response

            # Отправляем файл
            return send_file(
                temp_path,
                mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                as_attachment=True,
                download_name=f"seanses_{date_from[:10]}_{date_to[:10]}.docx",
            )
        except Exception as e:
            import logging

            logging.getLogger("web_portal.sessions").error(f"Error exporting word: {e}")
            return jsonify({"ok": False, "error": str(e)}), 500

    @app.post("/api/sessions/export-excel")
    @login_required
    def api_sessions_export_excel():
        """Выгрузка сеансов в Excel в формате: Время выхода, Частота, ID корреспондента, Группа, AES ключа, Color Voice, Время выхода (сек)."""
        try:
            from openpyxl.styles import Font, Alignment

            data = request.get_json(silent=True) or {}
            window = build_sessions_export_datetime_window(
                date_from=data.get("date_from", "").strip(),
                time_from=data.get("time_from", "00:00").strip(),
                date_to=data.get("date_to", "").strip(),
                time_to=data.get("time_to", "23:59").strip(),
            )
            datetime_from = window.datetime_from
            datetime_to = window.datetime_to
            date_from = window.date_from
            date_to = window.date_to

            p = db_path(DEFAULT_DB_NAME)
            ensure_db(p)
            seans_conn = get_request_seans_db()
            position_name = get_selected_position()
            if current_user.role == "admin" and not position_name:
                rows = get_seanses_rows_export(seans_conn, datetime_from, datetime_to)
            else:
                if not position_name:
                    return (
                        jsonify({"ok": False, "error": "Позиция не выбрана (выберите позицию слева)"}),
                        400,
                    )
                if not require_position_role(position_name, "export"):
                    return jsonify({"ok": False, "error": "Нет прав на экспорт"}), 403
                rows = get_seanses_rows_export(
                    seans_conn, datetime_from, datetime_to, client_name=position_name
                )

            wb = Workbook()
            ws = wb.active
            ws.title = "Сеансы"
            headers = [
                "Время выхода",
                "Частота",
                "ID корреспондента",
                "Группа",
                "AES ключа",
                "Color Voice",
                "Время выхода (сек)",
            ]
            for col, h in enumerate(headers, start=1):
                cell = ws.cell(row=1, column=col, value=h)
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", wrap_text=True)
            for row_idx, r in enumerate(rows, start=2):
                ws.cell(row=row_idx, column=1, value=r.get("date_time") or "")
                ws.cell(row=row_idx, column=2, value=r.get("frequency") or "")
                ws.cell(row=row_idx, column=3, value=r.get("id") or "")
                ws.cell(row=row_idx, column=4, value=r.get("group") or "")
                ws.cell(row=row_idx, column=5, value=r.get("aes_key") or "")
                ws.cell(row=row_idx, column=6, value=r.get("color_voice") or "")
                t = r.get("time_seconds")
                ws.cell(row=row_idx, column=7, value=round(t, 3) if t is not None else "")
            for col in range(1, len(headers) + 1):
                ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 18

            buf = BytesIO()
            wb.save(buf)
            buf.seek(0)
            filename = f"seanses_{date_from[:10]}_{date_to[:10]}.xlsx"
            return send_file(
                buf,
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                as_attachment=True,
                download_name=filename,
            )
        except Exception as e:
            logging.getLogger("web_portal.sessions").error(f"Error exporting excel: {e}")
            return jsonify({"ok": False, "error": str(e)}), 500

    @app.get("/api/sessions/favorites")
    @login_required
    def api_sessions_get_favorites():
        """Получает избранное пользователя для сеансов. ?with_unit_names=1 — добавить unit_name для учёта интенсивности."""
        try:
            conn = get_request_main_db()
            favorites = assemble_sessions_favorites_list(
                conn,
                user_id=int(current_user.id),
                with_unit_names=bool(request.args.get("with_unit_names")),
                include_units=bool(request.args.get("include_units")),
            )
            return jsonify({"ok": True, "favorites": favorites})
        except Exception as e:
            import logging

            logging.getLogger("web_portal.sessions").error(
                f"Error getting favorites: {e}"
            )
            return jsonify({"ok": False, "error": str(e)}), 500

    @app.post("/api/sessions/favorites/add")
    @login_required
    def api_sessions_add_favorite():
        """Добавляет запись в избранное: по (frequency, group) или по unit_name (для учёта интенсивности)."""
        try:
            data = request.get_json(silent=True) or {}
            p = db_path(DEFAULT_DB_NAME)
            ensure_db(p)
            conn = connect(p)
            try:
                payload = execute_sessions_add_favorite(
                    conn,
                    user_id=int(current_user.id),
                    frequency=data.get("frequency") or "",
                    group=data.get("group") or "",
                    unit_name=data.get("unit_name") or "",
                )
                return jsonify(payload)
            finally:
                conn.close()
        except SessionsUseCaseHTTP as e:
            return jsonify(e.payload), e.status_code
        except Exception as e:
            import logging

            logging.getLogger("web_portal.sessions").error(
                f"Error adding favorite: {e}"
            )
            return jsonify({"ok": False, "error": str(e)}), 500

    @app.post("/api/sessions/favorites/remove")
    @login_required
    def api_sessions_remove_favorite():
        try:
            data = request.get_json(silent=True) or {}
            p = db_path(DEFAULT_DB_NAME)
            ensure_db(p)
            conn = connect(p)
            try:
                payload = execute_sessions_remove_favorite(
                    conn,
                    user_id=int(current_user.id),
                    frequency=data.get("frequency") or "",
                    group=data.get("group") or "",
                    unit_name=data.get("unit_name") or "",
                )
                return jsonify(payload)
            finally:
                conn.close()
        except SessionsUseCaseHTTP as e:
            return jsonify(e.payload), e.status_code
        except Exception as e:
            import logging

            logging.getLogger("web_portal.sessions").error(
                f"Error removing favorite: {e}"
            )
            return jsonify({"ok": False, "error": str(e)}), 500

    @app.get("/api/sessions/id-watch")
    @login_required
    def api_sessions_id_watch_list():
        """Глобальный список отслеживаемых ID корреспондентов (общий для хаба)."""
        try:
            from web_portal.lib.sessions_id_watch import list_watched_correspondent_ids

            conn = get_request_main_db()
            items = list_watched_correspondent_ids(conn)
            return jsonify({"ok": True, "items": items, "count": len(items)})
        except Exception as e:
            logging.getLogger("web_portal.sessions").error(
                "Error listing id watch: %s", e
            )
            return jsonify({"ok": False, "error": str(e)}), 500

    @app.post("/api/sessions/id-watch/add")
    @login_required
    def api_sessions_id_watch_add():
        try:
            from web_portal.lib.sessions_id_watch import (
                add_watched_correspondent_id,
                add_watched_correspondent_ids,
                parse_correspondent_id_list,
                scan_recent_seans_for_watched_ids,
            )

            data = request.get_json(silent=True) or {}
            raw_ids = str(data.get("ids") or "").strip()
            cid = str(data.get("correspondent_id") or data.get("id") or "").strip()
            conn = get_request_main_db()
            author = str(getattr(current_user, "username", "") or "")

            if raw_ids:
                parsed = parse_correspondent_id_list(raw_ids)
                if not parsed:
                    return jsonify({"ok": False, "error": "Не найдено корректных ID"}), 400
                added = add_watched_correspondent_ids(
                    conn, parsed, created_by=author
                )
                scan_ids = added or parsed
                alerts_created = scan_recent_seans_for_watched_ids(
                    correspondent_ids=scan_ids
                )
                return jsonify(
                    {
                        "ok": True,
                        "added": added,
                        "count": len(added),
                        "alerts_created": alerts_created,
                    }
                )

            if not cid:
                return jsonify({"ok": False, "error": "correspondent_id обязателен"}), 400
            added_one = add_watched_correspondent_id(conn, cid, created_by=author)
            if not added_one:
                return jsonify({"ok": False, "error": "Некорректный ID или уже в списке"}), 400
            alerts_created = scan_recent_seans_for_watched_ids(correspondent_ids=[cid])
            return jsonify(
                {
                    "ok": True,
                    "correspondent_id": cid,
                    "added": [cid],
                    "count": 1,
                    "alerts_created": alerts_created,
                }
            )
        except Exception as e:
            logging.getLogger("web_portal.sessions").error("Error add id watch: %s", e)
            return jsonify({"ok": False, "error": str(e)}), 500

    @app.post("/api/sessions/id-watch/remove")
    @login_required
    def api_sessions_id_watch_remove():
        try:
            from web_portal.lib.sessions_id_watch import remove_watched_correspondent_id

            data = request.get_json(silent=True) or {}
            cid = str(data.get("correspondent_id") or data.get("id") or "").strip()
            if not cid:
                return jsonify({"ok": False, "error": "correspondent_id обязателен"}), 400
            conn = get_request_main_db()
            removed = remove_watched_correspondent_id(conn, cid)
            if not removed:
                return jsonify({"ok": False, "error": "ID не найден в списке"}), 404
            return jsonify({"ok": True, "correspondent_id": cid})
        except Exception as e:
            logging.getLogger("web_portal.sessions").error(
                "Error remove id watch: %s", e
            )
            return jsonify({"ok": False, "error": str(e)}), 500

    @app.get("/api/sessions/id-watch/alerts")
    @login_required
    def api_sessions_id_watch_alerts():
        """Новые оповещения для всех клиентов (poll с параметром since)."""
        try:
            from web_portal.lib.sessions_id_watch import list_id_watch_alerts_since

            since_raw = request.args.get("since") or "0"
            try:
                since = int(since_raw)
            except Exception:
                since = 0
            payload = list_id_watch_alerts_since(since)
            return jsonify({"ok": True, **payload})
        except Exception as e:
            logging.getLogger("web_portal.sessions").error(
                "Error id watch alerts: %s", e
            )
            return jsonify({"ok": False, "error": str(e)}), 500

    @app.post("/api/sessions/update-unit-name")
    @login_required
    def api_sessions_update_unit_name():
        """Ручное обновление названия подразделения (unit.name) по (frequency, group_)"""
        try:
            can_edit = can_edit_session_unit_name()
            if not can_edit:
                return (
                    jsonify(
                        {"ok": False, "error": "Нет прав на изменение подразделения"}
                    ),
                    403,
                )

            data = request.get_json(silent=True) or {}
            frequency = str(data.get("frequency") or "").strip()
            group = str(data.get("group") or "").strip()
            name = str(data.get("name") or "").strip()

            if not frequency or not group:
                return (
                    jsonify({"ok": False, "error": "frequency и group обязательны"}),
                    400,
                )

            p = db_path(DEFAULT_DB_NAME)
            ensure_db(p)
            conn = connect(p)
            try:
                res = update_unit_name(conn, frequency, group, name)
                return jsonify({"ok": True, "result": res})
            finally:
                conn.close()
        except Exception as e:
            import logging

            logging.getLogger("web_portal.sessions").error(
                f"Error updating unit name: {e}"
            )
            return jsonify({"ok": False, "error": str(e)}), 500

    @app.get("/api/sessions/unit-debug")
    @login_required
    def api_sessions_unit_debug():
        frequency = str(request.args.get("frequency") or "").strip()
        group_ = str(request.args.get("group") or "").strip()
        if not frequency or not group_:
            return jsonify({"ok": False, "error": "frequency/group обязательны"}), 400
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            unit_rows = conn.execute(
                """
                SELECT frequency, group_, name, manual, updated_at
                FROM unit
                WHERE group_ = ?
                """,
                (group_,),
            ).fetchall()
        finally:
            conn.close()

        from web_portal.config import search_online_db_path

        search_p = search_online_db_path()
        ensure_db(search_p)
        search_conn = connect(search_p)
        try:
            freq_prefix = _normalize_frequency_key(frequency)
            gid_tokens = re.findall(r"\d+", group_)
            like_gid = f"%{gid_tokens[0]}%" if gid_tokens else f"%{group_}%"
            like_freq = f"{freq_prefix}%"
            os_rows = search_conn.execute(
                """
                SELECT frequency, group_id, note, updated_at
                FROM online_search
                WHERE (group_id LIKE ? OR note LIKE ?)
                  AND frequency LIKE ?
                ORDER BY updated_at DESC
                LIMIT 50
                """,
                (like_gid, like_gid, like_freq),
            ).fetchall()
        finally:
            search_conn.close()

        return jsonify(
            {
                "ok": True,
                "query": {"frequency": frequency, "group": group_},
                "unit_rows": [
                    {
                        "frequency": str(r[0] or ""),
                        "group": str(r[1] or ""),
                        "name": str(r[2] or ""),
                        "manual": int(r[3] or 0),
                        "updated_at": str(r[4] or ""),
                    }
                    for r in unit_rows
                ],
                "online_search_rows": [
                    {
                        "frequency": str(r[0] or ""),
                        "group_id": str(r[1] or ""),
                        "note": str(r[2] or ""),
                        "updated_at": str(r[3] or ""),
                    }
                    for r in os_rows
                ],
            }
        )

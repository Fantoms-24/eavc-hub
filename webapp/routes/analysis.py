"""Анализ радиосетей (/analysis, /api/analysis/*), извлечено из app.py дословно.

Общая job-машинерия create_app передаётся параметрами. ML-оркестрация
вынесена в application.analysis.ml_use_cases (lib.ml с graceful fallback).
"""

from __future__ import annotations

import json
import logging
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

from flask import jsonify, redirect, render_template, request, url_for
from flask_login import current_user, login_required
from openpyxl import load_workbook

from web_portal.app_perf import perf_route as _perf_route, perf_span as _perf_span
from web_portal.app_runtime_caches import (
    analysis_cache_get as _analysis_cache_get,
    analysis_cache_key as _analysis_cache_key,
    analysis_cache_set as _analysis_cache_set,
    assignments_cache_get as _assignments_cache_get,
    assignments_cache_key as _assignments_cache_key,
    assignments_cache_set as _assignments_cache_set,
    callsigns_cache_get as _callsigns_cache_get,
    callsigns_cache_key as _callsigns_cache_key,
    callsigns_cache_set as _callsigns_cache_set,
    units_tree_cache_get as _units_tree_cache_get,
    units_tree_cache_key as _units_tree_cache_key,
    units_tree_cache_set as _units_tree_cache_set,
)
from web_portal.config import DEFAULT_DB_NAME, db_path, search_online_db_path
from web_portal.lib.ai_reports import run_period_report, run_shift_summary
from web_portal.lib.db import (
    connect,
    count_seanses_by_code,
    count_seanses_by_code_for_unit,
    ensure_db,
    get_ai_job,
    get_intercept_session,
    list_intercept_callsigns,
    list_online_search_unit_avatar_files,
    list_units_tree,
    update_ai_job,
)
from web_portal.lib.export_job_builders import (
    build_analysis_callsigns_export as _build_analysis_callsigns_export,
    build_analysis_full_report_docx as _build_analysis_full_report_docx,
    build_analysis_network_summary_export as _build_analysis_network_summary_export,
)
from web_portal.lib.export_jobs import export_results_dir as _export_results_dir
from web_portal.lib.flask_db import get_request_main_db, get_request_seans_db
from web_portal.positions import get_selected_position, require_position_role
from web_portal.webapp.auth import require_tab
from web_portal.webapp.context import AppContext
from web_portal.webapp.httputils import json_etag_response as _json_etag_response
from web_portal.webapp.timeutils import _utc_now_str
from web_portal.application.analysis import (
    AnalysisUseCaseHTTP,
    assemble_ai_reports_list_payload,
    assemble_analysis_graph_compare_payload,
    assemble_analysis_graph_dynamic_live_payload,
    assemble_analysis_graph_dynamic_timeline_payload,
    assemble_analysis_graph_dynamic_window_payload,
    assemble_analysis_graph_payload,
    assemble_analysis_graph_timeline_payload,
    assemble_analysis_keys_check_payload,
    assemble_analysis_pair_stats_rows,
    assemble_analysis_unit_stats_rows,
    assemble_arm_task_settings,
    build_analysis_period_bounds,
    build_full_report_job_params,
    execute_ai_feedback,
    execute_analysis_assign,
    execute_analysis_callsign_tag_update,
    execute_analysis_keys_import,
    execute_analysis_network_order_save,
    execute_analysis_network_order_shared,
    execute_arm_task_apply_audio_tasks,
    execute_arm_task_settings_save,
    execute_arm_task_view,
    execute_ml_dataset_build,
    execute_ml_dataset_build_forecast,
    execute_ml_feedback,
    execute_ml_predict,
    execute_ml_predict_forecast,
    execute_ml_train,
    execute_ml_train_forecast,
    execute_network_summary,
    fmt_analysis_dt,
    load_analysis_assignments_for_pair,
    ml_get_active_model_for_task,
    ml_get_settings_map,
    ml_set_setting,
    ml_setting_key,
    msk_now_naive,
    parse_analysis_datetime,
    parse_analysis_keys_check_range,
    parse_network_summary_range,
    parse_optional_session_id,
    parse_period_report_range,
    refresh_stale_ai_job,
    resolve_full_report_period,
    resolve_graph_timeline_window,
    save_full_report_graph_png,
)

_log = logging.getLogger("web_portal.webapp.routes.analysis")


def register_analysis_routes(app, ctx: AppContext):
    """Регистрирует маршруты анализа через AppContext. Код перенесён из create_app дословно."""
    logger = _log

    def attach_graph_unit_avatars(payload: dict) -> dict:
        """Добавляет фото подразделения к узлам и связям графа.

        Узел-корреспондент хранит своё подразделение в ``unit_name``. У связи
        это поле описывает группу, в которой корреспонденты встретились. Один
        запрос к справочнику даёт фото для обеих сущностей и для static-, и
        для dynamic-режима графа.
        """
        nodes = payload.get("nodes") or []
        edges = payload.get("edges") or []
        unit_keys = sorted(
            {
                str(item.get("unit_name") or "").strip()
                for item in [*nodes, *edges]
                if isinstance(item, dict)
                and str(item.get("unit_name") or "").strip()
            }
        )
        if not unit_keys:
            return payload
        search_p = search_online_db_path()
        ensure_db(search_p)
        search_conn = connect(search_p)
        try:
            avatar_files = list_online_search_unit_avatar_files(search_conn, unit_keys)
        finally:
            search_conn.close()
        for item in [*nodes, *edges]:
            if not isinstance(item, dict):
                continue
            unit_key = str(item.get("unit_name") or "").strip()
            avatar_file = avatar_files.get(unit_key)
            item["avatar_url"] = (
                url_for("api_online_search_unit_avatar_file", filename=avatar_file)
                if avatar_file
                else ""
            )
        return payload

    @app.before_request
    def reject_graph_rendering_on_hub():
        """HUB не строит графы и не тратит ресурсы на тяжёлые запросы графа."""
        if ctx.is_sync_hub() and request.path.startswith("/api/analysis/graph"):
            return (
                jsonify(
                    {
                        "ok": False,
                        "error": "Граф радиосети доступен только на центральном SERVER.",
                        "server_only": True,
                    }
                ),
                409,
            )

    @app.get("/analysis")
    @app.get("/analysis/<module>")
    @login_required
    @require_tab("tab_analysis")
    def analysis(module: str | None = None):
        analysis_module = (module or "callsigns").strip().lower()
        allowed = (
            "callsigns",
            "network",
            "graph",
            "keys",
            "crypto",
            "blanks-search",
            "ai",
            "arm-task",
        )
        if analysis_module not in allowed:
            return redirect(url_for("analysis"))
        return render_template(
            "analysis.html",
            analysis_module=analysis_module,
            graph_server_available=not ctx.is_sync_hub(),
            graph_all_positions=(current_user.role == "admin"),
        )

    @app.post("/api/analysis/callsigns/export")
    @login_required
    def api_analysis_callsigns_export():
        """Deprecated: используйте POST /api/analysis/callsigns/export-job."""
        return (
            jsonify(
                {
                    "ok": False,
                    "error": "Экспорт переведён в фон. Используйте /api/analysis/callsigns/export-job",
                    "use_job": True,
                }
            ),
            410,
        )

    _fmt_dt = fmt_analysis_dt
    _msk_now_naive = msk_now_naive
    _parse_dt = parse_analysis_datetime

    def _analysis_period_bounds_from_request() -> (
        tuple[datetime, datetime, dict[str, object]] | None
    ):
        bounds = build_analysis_period_bounds(
            period_start=(request.args.get("period_start") or "").strip(),
            period_end=(request.args.get("period_end") or "").strip(),
            days=str(request.args.get("days") or "1"),
        )
        if bounds is None:
            return None
        return bounds.start_naive, bounds.end_naive, bounds.meta

    def _analysis_ai_position() -> tuple[str | None, tuple | None]:
        pos = get_selected_position()
        if not pos and current_user.role == "admin":
            return "__all__", None
        if not pos:
            return None, (jsonify({"ok": False, "error": "Не выбрана позиция"}), 400)
        if not require_position_role(pos, "view"):
            return None, (jsonify({"ok": False, "error": "Нет доступа к позиции"}), 403)
        return pos, None

    def _build_analysis_callsigns_export_result(job_id: int, params: dict[str, object]) -> dict[str, object]:
        return _build_analysis_callsigns_export(ctx.base_dir, job_id, params)

    def _build_analysis_network_summary_export_result(
        job_id: int, params: dict[str, object]
    ) -> dict[str, object]:
        return _build_analysis_network_summary_export(ctx.base_dir, job_id, params)

    def _build_analysis_full_report_export_result(
        job_id: int, params: dict[str, object]
    ) -> dict[str, object]:
        return _build_analysis_full_report_docx(ctx.base_dir, job_id, params)

    @app.post("/api/analysis/callsigns/export-job")
    @_perf_route("api/analysis/callsigns/export-job")
    @login_required
    def api_analysis_callsigns_export_job():
        data = request.get_json(silent=True) or {}
        days = int(str(data.get("days") or "7").strip() or "7")
        if days not in {1, 3, 7, 30}:
            days = 7
        pos = get_selected_position()
        if not pos:
            return jsonify({"ok": False, "error": "Позиция не выбрана"}), 400
        if not require_position_role(pos, "view"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403
        params = {"position_name": str(pos), "days": days}
        job_id = ctx.create_export_job_record(
            job_type="analysis_callsigns",
            position_name=str(pos),
            params=params,
            created_by=ctx.ai_created_by(),
        )

        def _build() -> dict[str, object]:
            return _build_analysis_callsigns_export_result(job_id, params)

        payload, status = ctx.submit_export_job(
            job_id,
            "analysis-callsigns-export",
            _build,
        )
        return jsonify(payload), status

    @app.post("/api/analysis/network-summary/export-job")
    @_perf_route("api/analysis/network-summary/export-job")
    @login_required
    def api_analysis_network_summary_export_job():
        data = request.get_json(silent=True) or {}
        pos = get_selected_position()
        position_key = "__all__" if (current_user.role == "admin" and not pos) else str(pos or "")
        if position_key != "__all__":
            if not position_key:
                return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
            if not require_position_role(position_key, "view"):
                return jsonify({"ok": False, "error": "Нет прав"}), 403
        start_s = str(data.get("start") or "").strip()
        end_s = str(data.get("end") or "").strip()
        if not start_s or not end_s:
            return jsonify({"ok": False, "error": "start/end обязательны"}), 400
        params = {
            "position_name": position_key,
            "user_id": int(current_user.id),
            "start": start_s,
            "end": end_s,
            "cluster_order": data.get("cluster_order"),
            "custom_units": data.get("custom_units") or "",
            "favorite_units": data.get("favorite_units") or [],
            "favorites_only": bool(data.get("favorites_only")),
            "units_only": data.get("units_only") or data.get("unit_name") or "",
        }
        job_id = ctx.create_export_job_record(
            job_type="analysis_network_summary",
            position_name=position_key,
            params=params,
            created_by=ctx.ai_created_by(),
        )

        def _build() -> dict[str, object]:
            return _build_analysis_network_summary_export_result(job_id, params)

        payload, status = ctx.submit_export_job(
            job_id,
            "analysis-network-summary-export",
            _build,
        )
        return jsonify(payload), status

    @app.post("/api/analysis/full-report/export-job")
    @_perf_route("api/analysis/full-report/export-job")
    @login_required
    def api_analysis_full_report_export_job():
        data = request.get_json(silent=True) or {}
        pos = get_selected_position()
        position_key = "__all__" if (current_user.role == "admin" and not pos) else str(pos or "")
        if position_key != "__all__":
            if not position_key:
                return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
            if not require_position_role(position_key, "view"):
                return jsonify({"ok": False, "error": "Нет прав"}), 403

        try:
            period = resolve_full_report_period(data)
        except AnalysisUseCaseHTTP as e:
            return jsonify(e.payload), e.status_code

        params = build_full_report_job_params(
            data,
            position_key=position_key,
            user_id=int(current_user.id),
            period=period,
        )
        job_id = ctx.create_export_job_record(
            job_type="analysis_full_report",
            position_name=position_key,
            params=params,
            created_by=ctx.ai_created_by(),
        )

        png_path = save_full_report_graph_png(
            str(data.get("graph_png_base64") or ""),
            export_dir=_export_results_dir(ctx.base_dir),
            job_id=job_id,
        )
        if png_path:
            params["graph_png_path"] = png_path

        def _build() -> dict[str, object]:
            return _build_analysis_full_report_export_result(job_id, params)

        payload, status = ctx.submit_export_job(
            job_id,
            "analysis-full-report-export",
            _build,
        )
        return jsonify(payload), status

    def _can_view_ai_position(position_name: str) -> bool:
        pos = str(position_name or "").strip()
        if pos == "__all__":
            return bool(current_user.role == "admin")
        if not pos:
            return False
        return bool(require_position_role(pos, "view"))

    def _can_view_ai_report(conn, report_id: int) -> bool:
        row = conn.execute(
            "SELECT position_name FROM ai_reports WHERE id=?",
            (int(report_id),),
        ).fetchone()
        if not row:
            return False
        return _can_view_ai_position(str(row["position_name"] or ""))

    @app.post("/api/analysis/ai/shift-summary")
    @_perf_route("api/analysis/ai/shift-summary")
    @login_required
    def api_analysis_ai_shift_summary():
        """AI-сводка смены по бланкам перехватов."""
        pos, error_response = _analysis_ai_position()
        if error_response is not None:
            return error_response
        data = request.get_json(silent=True) or {}
        try:
            session_id = parse_optional_session_id(data.get("session_id"))
        except AnalysisUseCaseHTTP as e:
            return jsonify(e.payload), e.status_code

        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            if session_id is not None and pos != "__all__":
                sess = get_intercept_session(conn, session_id)
                if not sess:
                    return jsonify({"ok": False, "error": "Смена не найдена"}), 404
                sess_pos = str(sess.get("position_name") or "").strip()
                if sess_pos and sess_pos != pos:
                    return jsonify({"ok": False, "error": "Нет доступа к смене"}), 403
        except Exception as e:
            return jsonify({"ok": False, "error": str(e)}), 500
        finally:
            conn.close()

        created_by = ctx.ai_created_by()
        focus_query = str(data.get("focus_query") or "").strip()
        params = {"session_id": session_id, "focus_query": focus_query}
        job_id = ctx.create_ai_job_record(
            task_type="shift_summary",
            position_name=str(pos or ""),
            params=params,
            created_by=created_by,
        )

        def _run() -> None:
            bg_conn = None
            try:
                bg_conn = connect(p)
                with _perf_span("ai/shift-summary/run"):
                    run_shift_summary(
                        bg_conn,
                        position_name=str(pos or ""),
                        session_id=session_id,
                        focus_query=focus_query,
                        created_by=created_by,
                        job_id=job_id,
                    )
            except Exception as exc:
                fail_conn = connect(p)
                try:
                    update_ai_job(fail_conn, job_id, status="failed", error_text=str(exc))
                finally:
                    fail_conn.close()
                raise
            finally:
                if bg_conn is not None:
                    bg_conn.close()

        payload, status = ctx.submit_ai_job(job_id, "ai-shift-summary", _run)
        return jsonify(payload), status

    @app.post("/api/analysis/ai/period-report")
    @_perf_route("api/analysis/ai/period-report")
    @login_required
    def api_analysis_ai_period_report():
        """AI-анализ перехватов за период."""
        pos, error_response = _analysis_ai_position()
        if error_response is not None:
            return error_response
        data = request.get_json(silent=True) or {}
        try:
            start, end = parse_period_report_range(data)
        except AnalysisUseCaseHTTP as e:
            return jsonify(e.payload), e.status_code

        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        created_by = ctx.ai_created_by()
        job_id = ctx.create_ai_job_record(
            task_type="period_report",
            position_name=str(pos or ""),
            params={"start": start, "end": end},
            created_by=created_by,
        )

        def _run() -> None:
            bg_conn = None
            try:
                bg_conn = connect(p)
                with _perf_span("ai/period-report/run"):
                    run_period_report(
                        bg_conn,
                        position_name=str(pos or ""),
                        start_ts=start,
                        end_ts=end,
                        created_by=created_by,
                        job_id=job_id,
                    )
            except Exception as exc:
                fail_conn = connect(p)
                try:
                    update_ai_job(fail_conn, job_id, status="failed", error_text=str(exc))
                finally:
                    fail_conn.close()
                raise
            finally:
                if bg_conn is not None:
                    bg_conn.close()

        payload, status = ctx.submit_ai_job(job_id, "ai-period-report", _run)
        return jsonify(payload), status

    @app.get("/api/analysis/ai/jobs/<int:job_id>")
    @_perf_route("api/analysis/ai/job")
    @login_required
    def api_analysis_ai_job(job_id: int):
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            job = get_ai_job(conn, job_id)
            if not job:
                return jsonify({"ok": False, "error": "AI-задача не найдена"}), 404
            if not _can_view_ai_position(str(job.get("position_name") or "")):
                return jsonify({"ok": False, "error": "Нет доступа к AI-задаче"}), 403
            job = refresh_stale_ai_job(conn, job, job_id=job_id)
            return jsonify({"ok": True, "job": job})
        finally:
            conn.close()

    @app.get("/api/analysis/ai/reports")
    @login_required
    def api_analysis_ai_reports():
        pos, error_response = _analysis_ai_position()
        if error_response is not None:
            return error_response
        report_type = str(request.args.get("type") or "").strip()
        limit = request.args.get("limit", 30, type=int)
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            return jsonify(
                assemble_ai_reports_list_payload(
                    conn,
                    position_name=str(pos or ""),
                    report_type=report_type,
                    limit=limit,
                )
            )
        finally:
            conn.close()

    @app.post("/api/analysis/ai/feedback")
    @login_required
    def api_analysis_ai_feedback():
        data = request.get_json(silent=True) or {}
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            try:
                rid = int(data.get("report_id"))
            except Exception:
                return jsonify({"ok": False, "error": "Некорректный report_id"}), 400
            return jsonify(
                execute_ai_feedback(
                    conn,
                    report_id=rid,
                    rating=str(data.get("rating") or ""),
                    corrected_text=str(data.get("corrected_text") or ""),
                    comment=str(data.get("comment") or ""),
                    created_by=ctx.ai_created_by(),
                    can_view_report=_can_view_ai_report(conn, rid),
                )
            )
        except AnalysisUseCaseHTTP as e:
            return jsonify(e.payload), e.status_code
        finally:
            conn.close()

    @app.get("/api/analysis/state")
    @_perf_route("api/analysis/state")
    @login_required
    def api_analysis_state():
        pos = get_selected_position()
        if current_user.role == "admin" and not pos:
            return jsonify({"ok": True, "position": None, "units": [], "callsigns": []})
        if not pos:
            return jsonify({"ok": True, "position": None, "units": [], "callsigns": []})
        if not require_position_role(pos, "view"):
            return jsonify({"ok": False, "error": "Нет доступа к позиции"}), 403

        cs_key = _callsigns_cache_key(pos)
        ut_key = _units_tree_cache_key(pos)
        callsigns = _callsigns_cache_get(cs_key)
        units = _units_tree_cache_get(ut_key)
        if callsigns is not None and units is not None:
            return jsonify(
                {"ok": True, "position": pos, "units": units, "callsigns": callsigns}
            )

        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            if callsigns is None:
                callsigns = list_intercept_callsigns(conn, pos)
                _callsigns_cache_set(cs_key, callsigns)
            if units is None:
                units = list_units_tree(conn, position_name=pos)
                _units_tree_cache_set(ut_key, units)
            return jsonify(
                {"ok": True, "position": pos, "units": units, "callsigns": callsigns}
            )
        finally:
            conn.close()

    @app.get("/api/analysis/stats")
    @_perf_route("api/analysis/stats")
    @login_required
    def api_analysis_stats():
        pos = get_selected_position()
        if not pos:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if not require_position_role(pos, "view"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403
        frequency = str(request.args.get("frequency") or "").strip()
        group_code = str(request.args.get("group") or "").strip()
        if not frequency or not group_code:
            return jsonify({"ok": False, "error": "frequency и group обязательны"}), 400

        bounds = _analysis_period_bounds_from_request()
        if bounds is None:
            return jsonify(
                {
                    "ok": False,
                    "error": "Неверные period_start / period_end (формат YYYY-MM-DD, не более 400 суток).",
                }
            ), 400
        start_naive, end_naive, meta = bounds
        start_s = _fmt_dt(start_naive)
        end_s = _fmt_dt(end_naive)
        days_field = (
            int(meta["preset_days"])
            if meta.get("period_mode") == "rolling"
            else int(meta.get("days_metric") or 1)
        )

        cs_key = _callsigns_cache_key(pos)
        assign_key = _assignments_cache_key(pos, frequency, group_code)
        callsigns = _callsigns_cache_get(cs_key)
        assignments = _assignments_cache_get(assign_key)

        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        seans_conn = get_request_seans_db()
        try:
            counts = count_seanses_by_code(
                conn,
                position_name=pos,
                frequency=frequency,
                group_code=group_code,
                start_dt=start_s,
                end_dt=end_s,
                seans_conn=seans_conn,
            )
            if callsigns is None:
                callsigns = list_intercept_callsigns(conn, pos)
                _callsigns_cache_set(cs_key, callsigns)
            if assignments is None:
                assignments = load_analysis_assignments_for_pair(
                    conn,
                    position_name=pos,
                    frequency=frequency,
                    group_code=group_code,
                )
                _assignments_cache_set(assign_key, assignments)
        finally:
            conn.close()

        rows = assemble_analysis_pair_stats_rows(
            counts,
            callsigns,
            frequency=frequency,
            group_code=group_code,
        )
        return jsonify(
            {
                "ok": True,
                "position": pos,
                "frequency": frequency,
                "group": group_code,
                "days": days_field,
                "period_mode": str(meta.get("period_mode") or ""),
                "period_start_day": meta.get("period_start_day"),
                "period_end_day": meta.get("period_end_day"),
                "period_key": str(meta.get("period_key") or ""),
                "start": start_s,
                "end": end_s,
                "rows": rows,
                "assignments": assignments,
            }
        )

    @app.get("/api/analysis/arm-task/settings")
    @login_required
    def api_analysis_arm_task_settings():
        if current_user.role != "admin":
            return jsonify({"ok": False, "error": "Только администратор"}), 403
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            return jsonify(assemble_arm_task_settings(conn))
        finally:
            conn.close()

    @app.post("/api/analysis/arm-task/settings")
    @login_required
    def api_analysis_arm_task_settings_post():
        if current_user.role != "admin":
            return jsonify({"ok": False, "error": "Только администратор"}), 403
        data = request.get_json(silent=True) or {}
        dir_s = str(data.get("dir") or data.get("path") or "").strip()
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            return jsonify(execute_arm_task_settings_save(conn, dir_s))
        except AnalysisUseCaseHTTP as e:
            return jsonify(e.payload), e.status_code
        finally:
            conn.close()

    @app.get("/api/analysis/arm-task")
    @login_required
    def api_analysis_arm_task():
        pos = get_selected_position()
        if not pos:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if not require_position_role(pos, "view"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            return jsonify(execute_arm_task_view(conn, pos))
        except AnalysisUseCaseHTTP as e:
            return jsonify(e.payload), e.status_code
        except Exception as e:
            logger.exception("api_analysis_arm_task")
            return jsonify({"ok": False, "error": str(e)}), 500
        finally:
            conn.close()

    @app.post("/api/analysis/arm-task/apply-audio-tasks")
    @login_required
    def api_analysis_arm_task_apply_audio_tasks():
        pos = get_selected_position()
        if not pos:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if not require_position_role(pos, "view"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            return jsonify(
                execute_arm_task_apply_audio_tasks(
                    conn,
                    position_name=pos,
                    created_by=str(getattr(current_user, "username", "") or ""),
                )
            )
        except AnalysisUseCaseHTTP as e:
            return jsonify(e.payload), e.status_code
        except Exception as e:
            logger.exception("api_analysis_arm_task_apply_audio_tasks")
            return jsonify({"ok": False, "error": str(e)}), 500
        finally:
            conn.close()

    @app.get("/api/analysis/unit-stats")
    @_perf_route("api/analysis/unit-stats")
    @login_required
    def api_analysis_unit_stats():
        pos = get_selected_position()
        if not pos:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if not require_position_role(pos, "view"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403
        unit_name = str(request.args.get("unit_name") or "").strip()
        if not unit_name:
            return jsonify({"ok": False, "error": "unit_name обязателен"}), 400

        bounds = _analysis_period_bounds_from_request()
        if bounds is None:
            return jsonify(
                {
                    "ok": False,
                    "error": "Неверные period_start / period_end (формат YYYY-MM-DD, не более 400 суток).",
                }
            ), 400
        start_naive, end_naive, meta = bounds
        start_s = _fmt_dt(start_naive)
        end_s = _fmt_dt(end_naive)
        days_field = (
            int(meta["preset_days"])
            if meta.get("period_mode") == "rolling"
            else int(meta.get("days_metric") or 1)
        )

        cs_key = _callsigns_cache_key(pos)
        callsigns = _callsigns_cache_get(cs_key)

        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        seans_conn = get_request_seans_db()
        try:
            counts = count_seanses_by_code_for_unit(
                conn,
                position_name=pos,
                unit_name=unit_name,
                start_dt=start_s,
                end_dt=end_s,
                seans_conn=seans_conn,
            )
            if callsigns is None:
                callsigns = list_intercept_callsigns(conn, pos)
                _callsigns_cache_set(cs_key, callsigns)
        finally:
            conn.close()

        rows = assemble_analysis_unit_stats_rows(counts, callsigns)
        return jsonify(
            {
                "ok": True,
                "position": pos,
                "unit_name": unit_name,
                "days": days_field,
                "period_mode": str(meta.get("period_mode") or ""),
                "period_start_day": meta.get("period_start_day"),
                "period_end_day": meta.get("period_end_day"),
                "period_key": str(meta.get("period_key") or ""),
                "start": start_s,
                "end": end_s,
                "rows": rows,
            }
        )

    @app.post("/api/analysis/keys/import")
    @login_required
    def api_analysis_keys_import():
        file = request.files.get("file")
        if not file:
            return jsonify({"ok": False, "error": "Файл не выбран"}), 400

        wb = load_workbook(file, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))

        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            payload = execute_analysis_keys_import(
                conn, rows, now_ts=_utc_now_str()
            )
            return jsonify(payload)
        except AnalysisUseCaseHTTP as e:
            return jsonify(e.payload), int(e.status_code)
        finally:
            conn.close()

    @app.get("/api/analysis/keys/check")
    @login_required
    def api_analysis_keys_check():
        pos = get_selected_position()
        position_filter = None if (current_user.role == "admin" and not pos) else pos
        if not position_filter and current_user.role != "admin":
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if position_filter and not require_position_role(position_filter, "view"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403

        match_id = str(request.args.get("match_id") or "0").strip() == "1"
        debug = str(request.args.get("debug") or "").strip() in {"1", "true", "yes"}
        all_keys = str(request.args.get("all_keys") or "").strip() in {
            "1",
            "true",
            "yes",
        }
        start_s = str(request.args.get("start") or "").strip()
        end_s = str(request.args.get("end") or "").strip()
        start = None
        end = None
        try:
            parsed_range = parse_analysis_keys_check_range(start_s, end_s)
            if parsed_range is not None:
                start, end = parsed_range
        except AnalysisUseCaseHTTP as e:
            return jsonify(e.payload), int(e.status_code)
        new_since = str(request.args.get("new_since") or "").strip()

        main_conn = get_request_main_db()
        seans_conn = get_request_seans_db()
        payload = assemble_analysis_keys_check_payload(
            main_conn,
            seans_conn,
            position_filter=position_filter,
            match_id=match_id,
            debug=debug,
            all_keys=all_keys,
            start=start,
            end=end,
            new_since=new_since,
        )
        return jsonify(payload)

    @app.get("/api/analysis/graph")
    @_perf_route("api/analysis/graph")
    @login_required
    def api_analysis_graph():
        import datetime as _dt

        # Граф — единый серверный обзор. Администратор всегда видит все
        # источники, даже если в шапке интерфейса выбрана одна позиция.
        is_admin = current_user.role == "admin"
        pos = None if is_admin else get_selected_position()
        all_positions = is_admin
        if not pos and not all_positions:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if pos and not require_position_role(pos, "view"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403

        unit_name = str(request.args.get("unit_name") or "").strip()
        frequency = str(request.args.get("frequency") or "").strip()
        group_code = str(request.args.get("group") or "").strip()
        unit_query = str(request.args.get("unit_query") or "").strip()
        group_query = str(request.args.get("group_query") or "").strip()
        include_other = str(
            request.args.get("include_other") or ""
        ).strip().lower() in {
            "1",
            "true",
            "yes",
        }
        minutes = request.args.get("minutes", 60, type=int)
        minutes = max(5, min(720, minutes))
        max_nodes_mode = (
            str(request.args.get("max_nodes_mode") or "balanced").strip().lower()
        )
        if max_nodes_mode not in {"overview", "balanced", "dense"}:
            max_nodes_mode = "balanced"
        mode_defaults = {
            "overview": (120, 420),
            "balanced": (200, 800),
            "dense": (320, 1400),
        }
        default_nodes, default_edges = mode_defaults[max_nodes_mode]
        max_nodes_arg = request.args.get("max_nodes")
        max_edges_arg = request.args.get("max_edges")
        max_nodes = int(max_nodes_arg or default_nodes)
        max_nodes = max(50, min(500, max_nodes))
        max_edges = int(max_edges_arg or default_edges)
        max_edges = max(100, min(2500, max_edges))
        min_weight = request.args.get("min_weight", 1, type=int)
        min_weight = max(1, min(100, min_weight))
        cluster_by = str(request.args.get("cluster_by") or "unit").strip().lower()
        if cluster_by not in {"unit", "group"}:
            cluster_by = "unit"
        layout_hint = str(request.args.get("layout_hint") or "cose").strip().lower()
        focus_id = str(request.args.get("focus_id") or "").strip()

        start_override = _parse_dt(str(request.args.get("start") or ""))
        end_override = _parse_dt(str(request.args.get("end") or ""))
        days_arg = str(request.args.get("days") or "").strip()
        if start_override and end_override and end_override > start_override:
            start = start_override
            end = end_override
        elif days_arg:
            try:
                days = max(1, min(30, int(days_arg)))
            except Exception:
                days = 1
            end = _msk_now_naive()
            start = end - __import__("datetime").timedelta(days=days)
        else:
            end = _msk_now_naive()
            start = end - __import__("datetime").timedelta(minutes=minutes)
        start_s = _fmt_dt(start)
        end_s = _fmt_dt(end)
        cache_key = _analysis_cache_key("graph", pos=pos or "__all__")
        cached_payload = _analysis_cache_get(cache_key)
        if cached_payload is not None:
            return _json_etag_response(attach_graph_unit_avatars(cached_payload), max_age=30)

        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        seans_conn = get_request_seans_db()
        try:
            max_unit_rows = max(
                500, min(int(request.args.get("max_unit_rows") or 20000), 50000)
            )
            payload = assemble_analysis_graph_payload(
                conn,
                seans_conn,
                pos=pos,
                all_positions=all_positions,
                unit_name=unit_name,
                frequency=frequency,
                group_code=group_code,
                unit_query=unit_query,
                group_query=group_query,
                include_other=include_other,
                minutes=minutes,
                start=start,
                end=end,
                start_s=start_s,
                end_s=end_s,
                max_nodes=max_nodes,
                max_edges=max_edges,
                min_weight=min_weight,
                cluster_by=cluster_by,
                layout_hint=layout_hint,
                focus_id=focus_id,
                max_nodes_mode=max_nodes_mode,
                max_unit_rows=max_unit_rows,
            )
            attach_graph_unit_avatars(payload)
        finally:
            conn.close()

        historical_ttl = 10 * 60 if end < (_dt.datetime.now() - _dt.timedelta(minutes=5)) else 45
        _analysis_cache_set(cache_key, payload, ttl_sec=historical_ttl)
        return _json_etag_response(payload, max_age=30)

    @app.get("/api/analysis/graph-timeline")
    @_perf_route("api/analysis/graph-timeline")
    @login_required
    def api_analysis_graph_timeline():
        import datetime as _dt

        pos = get_selected_position()
        all_positions = current_user.role == "admin" and not pos
        if not pos and not all_positions:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if pos and not require_position_role(pos, "view"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403

        unit_name = str(request.args.get("unit_name") or "").strip()
        frequency = str(request.args.get("frequency") or "").strip()
        group_code = str(request.args.get("group") or "").strip()
        unit_query = str(request.args.get("unit_query") or "").strip()
        group_query = str(request.args.get("group_query") or "").strip()
        include_other = str(
            request.args.get("include_other") or ""
        ).strip().lower() in {"1", "true", "yes"}
        days = max(1, min(30, request.args.get("days", 1, type=int)))
        start_override = _parse_dt(str(request.args.get("start") or ""))
        end_override = _parse_dt(str(request.args.get("end") or ""))
        start, end, bucket_minutes = resolve_graph_timeline_window(
            days=days,
            start_override=start_override,
            end_override=end_override,
            now=_msk_now_naive(),
        )
        start_s = _fmt_dt(start)
        end_s = _fmt_dt(end)
        cache_key = _analysis_cache_key("graph-timeline", pos=pos or "__all__")
        cached_payload = _analysis_cache_get(cache_key)
        if cached_payload is not None:
            return _json_etag_response(cached_payload, max_age=30)

        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        seans_conn = get_request_seans_db()
        try:
            payload = assemble_analysis_graph_timeline_payload(
                conn,
                seans_conn,
                pos=pos,
                all_positions=all_positions,
                unit_name=unit_name,
                frequency=frequency,
                group_code=group_code,
                unit_query=unit_query,
                group_query=group_query,
                include_other=include_other,
                days=days,
                start=start,
                end=end,
                bucket_minutes=bucket_minutes,
                start_s=start_s,
                end_s=end_s,
            )
        finally:
            conn.close()

        historical_ttl = 10 * 60 if end < (_dt.datetime.now() - _dt.timedelta(minutes=5)) else 45
        _analysis_cache_set(cache_key, payload, ttl_sec=historical_ttl)
        return _json_etag_response(payload, max_age=30)

    @app.get("/api/analysis/graph-compare")
    @_perf_route("api/analysis/graph-compare")
    @login_required
    def api_analysis_graph_compare():
        pos = get_selected_position()
        all_positions = current_user.role == "admin" and not pos
        if not pos and not all_positions:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if pos and not require_position_role(pos, "view"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403

        unit_name = str(request.args.get("unit_name") or "").strip()
        frequency = str(request.args.get("frequency") or "").strip()
        group_code = str(request.args.get("group") or "").strip()
        unit_query = str(request.args.get("unit_query") or "").strip()
        group_query = str(request.args.get("group_query") or "").strip()
        include_other = str(
            request.args.get("include_other") or ""
        ).strip().lower() in {"1", "true", "yes"}
        days = max(1, min(30, request.args.get("days", 1, type=int)))
        start_override = _parse_dt(str(request.args.get("start") or ""))
        end_override = _parse_dt(str(request.args.get("end") or ""))
        start, end, bucket_minutes = resolve_graph_timeline_window(
            days=days,
            start_override=start_override,
            end_override=end_override,
            now=_msk_now_naive(),
        )
        prev_end = start
        prev_start = start - (end - start)
        start_s = _fmt_dt(start)
        end_s = _fmt_dt(end)
        prev_start_s = _fmt_dt(prev_start)
        prev_end_s = _fmt_dt(prev_end)
        cache_key = _analysis_cache_key("graph-compare", pos=pos or "__all__")
        cached_payload = _analysis_cache_get(cache_key)
        if cached_payload is not None:
            return jsonify(cached_payload)

        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        seans_conn = get_request_seans_db()
        try:
            payload = assemble_analysis_graph_compare_payload(
                conn,
                seans_conn,
                pos=pos,
                all_positions=all_positions,
                unit_name=unit_name,
                frequency=frequency,
                group_code=group_code,
                unit_query=unit_query,
                group_query=group_query,
                include_other=include_other,
                days=days,
                start=start,
                end=end,
                bucket_minutes=bucket_minutes,
                start_s=start_s,
                end_s=end_s,
                prev_start=prev_start,
                prev_end=prev_end,
                prev_start_s=prev_start_s,
                prev_end_s=prev_end_s,
            )
        finally:
            conn.close()

        _analysis_cache_set(cache_key, payload)
        return _json_etag_response(payload, max_age=30)

    @app.get("/api/analysis/graph-dynamic/timeline")
    @login_required
    def api_analysis_graph_dynamic_timeline():
        pos = get_selected_position()
        all_positions = current_user.role == "admin" and not pos
        if not pos and not all_positions:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if pos and not require_position_role(pos, "view"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403
        unit_name = str(request.args.get("unit_name") or "").strip()
        frequency = str(request.args.get("frequency") or "").strip()
        group_code = str(request.args.get("group") or "").strip()
        unit_query = str(request.args.get("unit_query") or "").strip()
        group_query = str(request.args.get("group_query") or "").strip()
        include_other = str(
            request.args.get("include_other") or ""
        ).strip().lower() in {"1", "true", "yes"}
        days = max(1, min(30, request.args.get("days", 1, type=int)))
        bucket_minutes = max(1, min(60, request.args.get("bucket_minutes", 5, type=int)))
        start_override = _parse_dt(str(request.args.get("start") or ""))
        end_override = _parse_dt(str(request.args.get("end") or ""))
        if start_override and end_override and end_override > start_override:
            start = start_override
            end = end_override
        else:
            end = _msk_now_naive()
            start = end - timedelta(days=days)
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        seans_conn = get_request_seans_db()
        try:
            return jsonify(
                assemble_analysis_graph_dynamic_timeline_payload(
                    conn,
                    seans_conn,
                    pos=pos,
                    all_positions=all_positions,
                    unit_name=unit_name,
                    frequency=frequency,
                    group_code=group_code,
                    unit_query=unit_query,
                    group_query=group_query,
                    include_other=include_other,
                    start=start,
                    end=end,
                    bucket_minutes=bucket_minutes,
                )
            )
        finally:
            conn.close()

    @app.get("/api/analysis/graph-dynamic/window")
    @login_required
    def api_analysis_graph_dynamic_window():
        pos = get_selected_position()
        all_positions = current_user.role == "admin" and not pos
        if not pos and not all_positions:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if pos and not require_position_role(pos, "view"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403
        unit_name = str(request.args.get("unit_name") or "").strip()
        frequency = str(request.args.get("frequency") or "").strip()
        group_code = str(request.args.get("group") or "").strip()
        unit_query = str(request.args.get("unit_query") or "").strip()
        group_query = str(request.args.get("group_query") or "").strip()
        include_other = str(
            request.args.get("include_other") or ""
        ).strip().lower() in {"1", "true", "yes"}
        cluster_by = str(request.args.get("cluster_by") or "unit").strip().lower()
        if cluster_by not in {"unit", "group"}:
            cluster_by = "unit"
        max_nodes_mode = (
            str(request.args.get("max_nodes_mode") or "balanced").strip().lower()
        )
        mode_defaults = {
            "overview": (120, 420),
            "balanced": (200, 800),
            "dense": (320, 1400),
        }
        default_nodes, default_edges = mode_defaults.get(
            max_nodes_mode, mode_defaults["balanced"]
        )
        max_nodes = max(
            50, min(500, request.args.get("max_nodes", default_nodes, type=int))
        )
        max_edges = max(
            100, min(2500, request.args.get("max_edges", default_edges, type=int))
        )
        min_weight = max(1, min(100, request.args.get("min_weight", 1, type=int)))
        window_minutes = max(
            5,
            min(
                720,
                int(
                    request.args.get("minutes")
                    or request.args.get("window_minutes")
                    or 60
                ),
            ),
        )
        cursor_ts = _parse_dt(str(request.args.get("cursor_ts") or ""))
        end_dt = cursor_ts or _msk_now_naive()
        start_dt = end_dt - timedelta(minutes=window_minutes)
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        seans_conn = get_request_seans_db()
        try:
            payload = assemble_analysis_graph_dynamic_window_payload(
                conn,
                seans_conn,
                pos=pos,
                all_positions=all_positions,
                unit_name=unit_name,
                frequency=frequency,
                group_code=group_code,
                unit_query=unit_query,
                group_query=group_query,
                include_other=include_other,
                cluster_by=cluster_by,
                min_weight=min_weight,
                max_nodes=max_nodes,
                max_edges=max_edges,
                window_minutes=window_minutes,
                start_dt=start_dt,
                end_dt=end_dt,
            )
            return jsonify(attach_graph_unit_avatars(payload))
        finally:
            conn.close()

    @app.get("/api/analysis/graph-dynamic/live")
    @login_required
    def api_analysis_graph_dynamic_live():
        pos = get_selected_position()
        all_positions = current_user.role == "admin" and not pos
        if not pos and not all_positions:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if pos and not require_position_role(pos, "view"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403
        unit_name = str(request.args.get("unit_name") or "").strip()
        frequency = str(request.args.get("frequency") or "").strip()
        group_code = str(request.args.get("group") or "").strip()
        unit_query = str(request.args.get("unit_query") or "").strip()
        group_query = str(request.args.get("group_query") or "").strip()
        include_other = str(
            request.args.get("include_other") or ""
        ).strip().lower() in {"1", "true", "yes"}
        bucket_minutes = max(1, min(60, request.args.get("bucket_minutes", 5, type=int)))
        since_dt = _parse_dt(str(request.args.get("since_ts") or ""))
        end_dt = _msk_now_naive()
        start_dt = since_dt or (end_dt - timedelta(minutes=max(10, bucket_minutes * 2)))
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        seans_conn = get_request_seans_db()
        try:
            return jsonify(
                assemble_analysis_graph_dynamic_live_payload(
                    conn,
                    seans_conn,
                    pos=pos,
                    unit_name=unit_name,
                    frequency=frequency,
                    group_code=group_code,
                    unit_query=unit_query,
                    group_query=group_query,
                    include_other=include_other,
                    bucket_minutes=bucket_minutes,
                    start_dt=start_dt,
                    end_dt=end_dt,
                )
            )
        finally:
            conn.close()

    @app.post("/api/analysis/assign")
    @login_required
    def api_analysis_assign():
        pos = get_selected_position()
        if not pos:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if not require_position_role(pos, "start"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403
        data = request.get_json(silent=True) or {}
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            payload = execute_analysis_assign(
                conn,
                position_name=pos,
                frequency=str(data.get("frequency") or ""),
                group_code=str(data.get("group") or ""),
                role_type=str(data.get("role_type") or ""),
                code=str(data.get("code") or ""),
                mode=str(data.get("mode") or ""),
                sync_hub=ctx.is_sync_hub(),
                now_ts=ctx.now_ts(),
            )
            return jsonify(payload)
        except AnalysisUseCaseHTTP as e:
            return jsonify(e.payload), int(e.status_code)
        finally:
            conn.close()

    @app.post("/api/analysis/callsign-tag")
    @login_required
    def api_analysis_callsign_tag():
        pos = get_selected_position()
        if not pos:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if not require_position_role(pos, "edit"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403
        data = request.get_json(silent=True) or {}
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            payload = execute_analysis_callsign_tag_update(
                conn,
                position_name=pos,
                callsign_id=int(data.get("callsign_id") or 0),
                tag=str(data.get("tag") or ""),
                tag_desc=str(data.get("tag_desc") or ""),
                tag_color=str(data.get("tag_color") or ""),
                sync_hub=ctx.is_sync_hub(),
            )
            return jsonify(payload)
        except AnalysisUseCaseHTTP as e:
            return jsonify(e.payload), int(e.status_code)
        finally:
            conn.close()

    @app.get("/api/analysis/ml/settings")
    @login_required
    def api_analysis_ml_settings_get():
        pos = get_selected_position()
        if not pos:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if not require_position_role(pos, "view"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            return jsonify(
                {"ok": True, "settings": ml_get_settings_map(conn, position_name=pos)}
            )
        finally:
            conn.close()

    @app.post("/api/analysis/ml/settings")
    @login_required
    def api_analysis_ml_settings_set():
        pos = get_selected_position()
        if not pos:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if not require_position_role(pos, "edit"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403
        payload = request.get_json(silent=True) or {}
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            window_minutes = max(5, min(720, int(payload.get("window_minutes") or 60)))
            train_epochs = max(5, min(1000, int(payload.get("train_epochs") or 60)))
            train_lr = max(0.00001, min(0.1, float(payload.get("train_lr") or 0.001)))
            auto_activate = bool(payload.get("auto_activate", True))
            retrain_min_samples = max(
                10, min(10000, int(payload.get("retrain_min_samples") or 30))
            )
            retrain_step = max(1, min(1000, int(payload.get("retrain_step") or 10)))
            ml_set_setting(
                conn,
                position_name=pos,
                name="window_minutes",
                value=str(window_minutes),
            )
            ml_set_setting(
                conn,
                position_name=pos,
                name="train_epochs",
                value=str(train_epochs),
            )
            ml_set_setting(
                conn,
                position_name=pos,
                name="train_lr",
                value=str(train_lr),
            )
            ml_set_setting(
                conn,
                position_name=pos,
                name="auto_activate",
                value="1" if auto_activate else "0",
            )
            ml_set_setting(
                conn,
                position_name=pos,
                name="retrain_min_samples",
                value=str(retrain_min_samples),
            )
            ml_set_setting(
                conn,
                position_name=pos,
                name="retrain_step",
                value=str(retrain_step),
            )
            conn.commit()
            return jsonify(
                {"ok": True, "settings": ml_get_settings_map(conn, position_name=pos)}
            )
        except Exception as e:
            return jsonify({"ok": False, "error": str(e)}), 400
        finally:
            conn.close()

    @app.post("/api/analysis/ml/event-tag")
    @login_required
    def api_analysis_ml_event_tag():
        pos = get_selected_position()
        if not pos:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if not require_position_role(pos, "edit"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403
        payload = request.get_json(silent=True) or {}
        start_ts = str(payload.get("start_ts") or "").strip()
        end_ts = str(payload.get("end_ts") or "").strip()
        label = str(payload.get("label") or "").strip()
        tag = str(payload.get("tag") or "").strip()
        description = str(payload.get("description") or "").strip()
        status = str(payload.get("status") or "validated").strip().lower()
        granularity = str(payload.get("granularity") or "window").strip().lower()
        if not start_ts or not end_ts or not label:
            return (
                jsonify({"ok": False, "error": "start_ts/end_ts/label обязательны"}),
                400,
            )
        try:
            s_dt = _parse_dt(start_ts)
            e_dt = _parse_dt(end_ts)
            if not s_dt or not e_dt or e_dt <= s_dt:
                return jsonify({"ok": False, "error": "Неверный интервал"}), 400
            if granularity == "minute":
                s_dt = s_dt.replace(second=0, microsecond=0)
                e_dt = e_dt.replace(second=0, microsecond=0)
                if e_dt <= s_dt:
                    e_dt = s_dt + timedelta(minutes=1)
        except Exception:
            return jsonify({"ok": False, "error": "Неверный формат даты"}), 400
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            conn.execute(
                """
                INSERT INTO ml_event_tags
                    (position_name, start_ts, end_ts, label, tag, description, status, validated_by, validated_at, meta_json, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP, '{}', CURRENT_TIMESTAMP)
                """,
                (
                    pos,
                    _fmt_dt(s_dt),
                    _fmt_dt(e_dt),
                    label,
                    tag,
                    description,
                    (
                        status
                        if status in {"validated", "rejected", "draft"}
                        else "validated"
                    ),
                    str(getattr(current_user, "username", "") or ""),
                ),
            )
            conn.execute(
                "UPDATE ml_event_tags SET meta_json=? WHERE id=last_insert_rowid()",
                (
                    json.dumps(
                        {
                            "granularity": (
                                granularity
                                if granularity in {"minute", "window"}
                                else "window"
                            )
                        },
                        ensure_ascii=False,
                    ),
                ),
            )
            conn.commit()
            return jsonify({"ok": True})
        except Exception as e:
            return jsonify({"ok": False, "error": str(e)}), 400
        finally:
            conn.close()

    @app.get("/api/analysis/ml/event-tags")
    @login_required
    def api_analysis_ml_event_tags():
        pos = get_selected_position()
        if not pos:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if not require_position_role(pos, "view"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403
        start = _parse_dt(str(request.args.get("start") or ""))
        end = _parse_dt(str(request.args.get("end") or ""))
        where = "position_name=?"
        params: list[object] = [pos]
        if start and end and end > start:
            where += " AND NOT (end_ts <= ? OR start_ts >= ?)"
            params.extend([_fmt_dt(start), _fmt_dt(end)])
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            rows = conn.execute(
                f"""
                SELECT id, start_ts, end_ts, label, tag, description, status, validated_by, validated_at, created_at
                FROM ml_event_tags
                WHERE {where}
                ORDER BY start_ts DESC
                LIMIT 500
                """,
                tuple(params),
            ).fetchall()
            return jsonify(
                {
                    "ok": True,
                    "items": [
                        {
                            "id": int(r["id"]),
                            "start_ts": str(r["start_ts"] or ""),
                            "end_ts": str(r["end_ts"] or ""),
                            "label": str(r["label"] or ""),
                            "tag": str(r["tag"] or ""),
                            "description": str(r["description"] or ""),
                            "status": str(r["status"] or ""),
                            "validated_by": str(r["validated_by"] or ""),
                            "validated_at": str(r["validated_at"] or ""),
                            "created_at": str(r["created_at"] or ""),
                        }
                        for r in (rows or [])
                    ],
                }
            )
        finally:
            conn.close()

    @app.post("/api/analysis/ml/event-tag/update")
    @login_required
    def api_analysis_ml_event_tag_update():
        pos = get_selected_position()
        if not pos:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if not require_position_role(pos, "edit"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403
        payload = request.get_json(silent=True) or {}
        tag_id = int(payload.get("id") or 0)
        if tag_id <= 0:
            return jsonify({"ok": False, "error": "id обязателен"}), 400
        label = str(payload.get("label") or "").strip()
        tag = str(payload.get("tag") or "").strip()
        description = str(payload.get("description") or "").strip()
        status = str(payload.get("status") or "").strip().lower()
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            row = conn.execute(
                "SELECT id FROM ml_event_tags WHERE id=? AND position_name=?",
                (tag_id, pos),
            ).fetchone()
            if not row:
                return jsonify({"ok": False, "error": "Тег не найден"}), 404
            updates: list[str] = []
            params: list[object] = []
            if label:
                updates.append("label=?")
                params.append(label)
            if tag:
                updates.append("tag=?")
                params.append(tag)
            if description:
                updates.append("description=?")
                params.append(description)
            if status in {"validated", "rejected", "draft"}:
                updates.append("status=?")
                params.append(status)
                updates.append("validated_by=?")
                params.append(str(getattr(current_user, "username", "") or ""))
                updates.append("validated_at=CURRENT_TIMESTAMP")
            if not updates:
                return jsonify({"ok": False, "error": "Нет полей для обновления"}), 400
            updates.append("updated_at=CURRENT_TIMESTAMP")
            conn.execute(
                f"UPDATE ml_event_tags SET {', '.join(updates)} WHERE id=? AND position_name=?",
                (*params, tag_id, pos),
            )
            conn.commit()
            return jsonify({"ok": True})
        except Exception as e:
            return jsonify({"ok": False, "error": str(e)}), 400
        finally:
            conn.close()

    @app.post("/api/analysis/ml/dataset-build")
    @login_required
    def api_analysis_ml_dataset_build():
        pos = get_selected_position()
        if not pos:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if not require_position_role(pos, "edit"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403
        payload = request.get_json(silent=True) or {}
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        seans_conn = get_request_seans_db()
        try:
            result = execute_ml_dataset_build(
                conn,
                seans_conn,
                position_name=pos,
                body=payload,
            )
            return jsonify(result)
        except Exception as e:
            return jsonify({"ok": False, "error": str(e)}), 400
        finally:
            conn.close()

    @app.post("/api/analysis/ml/train")
    @login_required
    def api_analysis_ml_train():
        pos = get_selected_position()
        if not pos:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if not require_position_role(pos, "edit"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403
        payload = request.get_json(silent=True) or {}
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            result = execute_ml_train(
                conn,
                position_name=pos,
                body=payload,
                activated_at=_utc_now_str(),
            )
            return jsonify(result)
        except Exception as e:
            return jsonify({"ok": False, "error": str(e)}), 400
        finally:
            conn.close()

    @app.post("/api/analysis/ml/dataset-build-forecast")
    @login_required
    def api_analysis_ml_dataset_build_forecast():
        pos = get_selected_position()
        if not pos:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if not require_position_role(pos, "edit"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403
        payload = request.get_json(silent=True) or {}
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        seans_conn = get_request_seans_db()
        try:
            result = execute_ml_dataset_build_forecast(
                conn,
                seans_conn,
                position_name=pos,
                body=payload,
            )
            return jsonify(result)
        except Exception as e:
            return jsonify({"ok": False, "error": str(e)}), 400
        finally:
            conn.close()

    @app.post("/api/analysis/ml/train-forecast")
    @login_required
    def api_analysis_ml_train_forecast():
        pos = get_selected_position()
        if not pos:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if not require_position_role(pos, "edit"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403
        payload = request.get_json(silent=True) or {}
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            result = execute_ml_train_forecast(
                conn,
                position_name=pos,
                body=payload,
                activated_at=_utc_now_str(),
            )
            return jsonify(result)
        except Exception as e:
            return jsonify({"ok": False, "error": str(e)}), 400
        finally:
            conn.close()

    @app.post("/api/analysis/ml/predict-forecast")
    @login_required
    def api_analysis_ml_predict_forecast():
        pos = get_selected_position()
        if not pos:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if not require_position_role(pos, "view"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403
        payload = request.get_json(silent=True) or {}
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        seans_conn = get_request_seans_db()
        try:
            result = execute_ml_predict_forecast(
                conn,
                seans_conn,
                position_name=pos,
                body=payload,
            )
            return jsonify(result)
        except RuntimeError as e:
            return jsonify({"ok": False, "error": str(e)}), 400
        except Exception as e:
            return jsonify({"ok": False, "error": str(e)}), 400
        finally:
            conn.close()

    @app.get("/api/analysis/ml/model-status")
    @login_required
    def api_analysis_ml_model_status():
        pos = get_selected_position()
        if not pos:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if not require_position_role(pos, "view"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            model = ml_get_active_model_for_task(conn, task_type="multiclass_event")
            forecast_models = conn.execute(
                """
                SELECT task_type, model_version, metrics_json, is_active
                FROM ml_models
                WHERE task_type LIKE 'forecast_h%'
                ORDER BY id DESC
                LIMIT 20
                """
            ).fetchall()
            all_rows = conn.execute(
                """
                SELECT model_version, metrics_json, is_active, created_at, activated_at
                FROM ml_models
                ORDER BY id DESC
                LIMIT 20
                """
            ).fetchall()
            pending_validated = int(
                (
                    conn.execute(
                        """
                        SELECT COUNT(*)
                        FROM ml_event_tags
                        WHERE position_name=? AND status='validated'
                        """,
                        (pos,),
                    ).fetchone()[0]
                )
                or 0
            )
            retrain_setting = conn.execute(
                "SELECT value FROM ml_settings WHERE key=?",
                (ml_setting_key(pos, "retrain_recommended"),),
            ).fetchone()
            retrain_recommended = (
                str((retrain_setting["value"] if retrain_setting else "0") or "0")
                == "1"
            )
            ml_settings = ml_get_settings_map(conn, position_name=pos)
            return jsonify(
                {
                    "ok": True,
                    "active_model": (
                        {
                            "model_version": str(model["model_version"] or ""),
                            "metrics": json.loads(str(model["metrics_json"] or "{}")),
                        }
                        if model
                        else None
                    ),
                    "models": [
                        {
                            "model_version": str(r["model_version"] or ""),
                            "is_active": int(r["is_active"] or 0) == 1,
                            "metrics": json.loads(str(r["metrics_json"] or "{}")),
                            "created_at": str(r["created_at"] or ""),
                            "activated_at": str(r["activated_at"] or ""),
                        }
                        for r in (all_rows or [])
                    ],
                    "diagnostics": {
                        "pending_validated_samples": pending_validated,
                        "retrain_recommended": retrain_recommended,
                        "fallback_enabled": True,
                    },
                    "forecast_models": [
                        {
                            "task_type": str(r["task_type"] or ""),
                            "model_version": str(r["model_version"] or ""),
                            "is_active": int(r["is_active"] or 0) == 1,
                            "metrics": json.loads(str(r["metrics_json"] or "{}")),
                        }
                        for r in (forecast_models or [])
                    ],
                    "settings": ml_settings,
                }
            )
        finally:
            conn.close()

    @app.post("/api/analysis/ml/model-activate")
    @login_required
    def api_analysis_ml_model_activate():
        pos = get_selected_position()
        if not pos:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if not require_position_role(pos, "edit"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403
        payload = request.get_json(silent=True) or {}
        model_version = str(payload.get("model_version") or "").strip()
        if not model_version:
            return jsonify({"ok": False, "error": "model_version обязателен"}), 400
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            row = conn.execute(
                "SELECT id, task_type, metrics_json FROM ml_models WHERE model_version=?",
                (model_version,),
            ).fetchone()
            if not row:
                return jsonify({"ok": False, "error": "Модель не найдена"}), 404
            task_type = str(row["task_type"] or "multiclass_event")
            candidate_metrics = json.loads(str(row["metrics_json"] or "{}"))
            candidate_f1 = float(candidate_metrics.get("val_f1_macro") or 0.0)
            active_row = conn.execute(
                """
                SELECT model_version, metrics_json
                FROM ml_models
                WHERE is_active=1 AND task_type=?
                ORDER BY id DESC
                LIMIT 1
                """,
                (task_type,),
            ).fetchone()
            if active_row:
                active_metrics = json.loads(str(active_row["metrics_json"] or "{}"))
                active_f1 = float(active_metrics.get("val_f1_macro") or 0.0)
                if candidate_f1 + 0.10 < active_f1:
                    return (
                        jsonify(
                            {
                                "ok": False,
                                "error": f"Качество ниже активной модели ({candidate_f1:.3f} < {active_f1:.3f} - 0.10).",
                            }
                        ),
                        400,
                    )
            if task_type.startswith("forecast_h"):
                conn.execute(
                    "UPDATE ml_models SET is_active=0 WHERE task_type=?", (task_type,)
                )
            else:
                conn.execute(
                    "UPDATE ml_models SET is_active=0 WHERE task_type NOT LIKE 'forecast_h%'"
                )
            conn.execute(
                "UPDATE ml_models SET is_active=1, activated_at=? WHERE model_version=?",
                (_utc_now_str(), model_version),
            )
            conn.commit()
            return jsonify({"ok": True})
        finally:
            conn.close()

    @app.post("/api/analysis/ml/predict")
    @login_required
    def api_analysis_ml_predict():
        pos = get_selected_position()
        if not pos:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if not require_position_role(pos, "view"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403
        payload = request.get_json(silent=True) or {}
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        seans_conn = get_request_seans_db()
        try:
            result = execute_ml_predict(
                conn,
                seans_conn,
                position_name=pos,
                body=payload,
            )
            return jsonify(result)
        except RuntimeError as e:
            return jsonify({"ok": False, "error": str(e)}), 400
        except Exception as e:
            return jsonify({"ok": False, "error": str(e)}), 400
        finally:
            conn.close()

    @app.get("/api/analysis/ml/explain")
    @login_required
    def api_analysis_ml_explain():
        pos = get_selected_position()
        if not pos:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if not require_position_role(pos, "view"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403
        prediction_id = request.args.get("prediction_id", 0, type=int)
        if prediction_id <= 0:
            return jsonify({"ok": False, "error": "prediction_id обязателен"}), 400
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            row = conn.execute(
                """
                SELECT id, start_ts, end_ts, top_label, top_probability, probabilities_json, reasons_json, hotspots_json, model_version
                FROM ml_predictions
                WHERE id=? AND position_name=?
                """,
                (prediction_id, pos),
            ).fetchone()
            if not row:
                return jsonify({"ok": False, "error": "Прогноз не найден"}), 404
            return jsonify(
                {
                    "ok": True,
                    "prediction_id": int(row["id"]),
                    "start_ts": str(row["start_ts"] or ""),
                    "end_ts": str(row["end_ts"] or ""),
                    "top_label": str(row["top_label"] or ""),
                    "top_probability": float(row["top_probability"] or 0.0),
                    "model_version": str(row["model_version"] or ""),
                    "probabilities": json.loads(str(row["probabilities_json"] or "[]")),
                    "reasons": json.loads(str(row["reasons_json"] or "[]")),
                    "hotspots": json.loads(str(row["hotspots_json"] or "[]")),
                }
            )
        finally:
            conn.close()

    @app.post("/api/analysis/ml/feedback")
    @login_required
    def api_analysis_ml_feedback():
        pos = get_selected_position()
        if not pos:
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if not require_position_role(pos, "edit"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403
        payload = request.get_json(silent=True) or {}
        prediction_id = int(payload.get("prediction_id") or 0)
        feedback_label = str(payload.get("feedback_label") or "").strip()
        if prediction_id <= 0 or not feedback_label:
            return (
                jsonify(
                    {"ok": False, "error": "prediction_id и feedback_label обязательны"}
                ),
                400,
            )
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            result = execute_ml_feedback(
                conn,
                position_name=pos,
                body=payload,
                reviewer=str(getattr(current_user, "username", "") or ""),
            )
            return jsonify(result)
        except LookupError as e:
            return jsonify({"ok": False, "error": str(e)}), 404
        except ValueError as e:
            return jsonify({"ok": False, "error": str(e)}), 400
        except Exception as e:
            return jsonify({"ok": False, "error": str(e)}), 400
        finally:
            conn.close()

    @app.get("/api/analysis/network-summary")
    @_perf_route("api/analysis/network-summary")
    @login_required
    def api_analysis_network_summary():
        """
        Общий анализ радиосети (кластеризация по unit_name):
        Возвращает данные за период и за предыдущий период такой же длительности.
        """
        pos = get_selected_position()
        position_filter = None if (current_user.role == "admin" and not pos) else pos
        position_key = position_filter or "__all__"
        if not position_filter and current_user.role != "admin":
            return jsonify({"ok": False, "error": "Не выбрана позиция"}), 400
        if position_filter and not require_position_role(position_filter, "view"):
            return jsonify({"ok": False, "error": "Нет прав"}), 403

        favorites_only = (
            str(request.args.get("favorites_only") or "").strip().lower()
            in {"1", "true", "yes", "on"}
        )
        favorite_units_raw = str(request.args.get("favorite_units") or "").strip()
        try:
            start, end = parse_network_summary_range(
                str(request.args.get("start") or "").strip(),
                str(request.args.get("end") or "").strip(),
            )
        except AnalysisUseCaseHTTP as e:
            return jsonify(e.payload), int(e.status_code)

        cache_key = _analysis_cache_key("network-summary", pos=position_filter)
        cached_payload = _analysis_cache_get(cache_key)
        if cached_payload is not None:
            return jsonify(cached_payload)

        try:
            main_conn = get_request_main_db()
            seans_conn = get_request_seans_db()
            payload = execute_network_summary(
                main_conn,
                seans_conn,
                position_filter=position_filter,
                position_key=position_key,
                all_positions=bool(current_user.role == "admin" and not pos),
                start=start,
                end=end,
                favorites_only=favorites_only,
                favorite_units_raw=favorite_units_raw,
                custom_units_raw=str(request.args.get("custom_units") or "").strip(),
                units_only_raw=str(
                    request.args.get("units_only") or request.args.get("unit_name") or ""
                ).strip(),
                user_id=int(current_user.id),
            )
            _analysis_cache_set(cache_key, payload)
            return jsonify(payload)
        except Exception as e:
            logger.exception("api_analysis_network_summary failed")
            return jsonify({"ok": False, "error": str(e)}), 500

    @app.post("/api/analysis/network-order/shared")
    @login_required
    def api_analysis_network_order_shared():
        can_manage = bool(
            current_user.role == "admin" or current_user.has_perm("manage_users")
        )
        if not can_manage:
            return jsonify({"ok": False, "error": "Нет прав"}), 403

        pos = get_selected_position()
        position_filter = None if (current_user.role == "admin" and not pos) else pos
        position_key = position_filter or "__all__"

        data = request.get_json(silent=True) or {}
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            payload = execute_analysis_network_order_shared(
                conn,
                position_key=position_key,
                order_raw=data.get("order") or [],
            )
            return jsonify(payload)
        except AnalysisUseCaseHTTP as e:
            return jsonify(e.payload), int(e.status_code)
        finally:
            conn.close()

    @app.post("/api/analysis/network-order/save")
    @login_required
    def api_analysis_network_order_save():
        """Сохраняет порядок колонок общего анализа радиосети в БД для текущей позиции."""
        pos = get_selected_position()
        position_key = str(pos or "").strip() or "__all__"
        if position_key == "__all__" and current_user.role != "admin":
            return jsonify(
                {"ok": False, "error": "Выберите позицию для сохранения порядка"}
            ), 400
        if position_key != "__all__":
            if not require_position_role(position_key, "view"):
                return jsonify({"ok": False, "error": "Нет прав"}), 403

        data = request.get_json(silent=True) or {}
        p = db_path(DEFAULT_DB_NAME)
        ensure_db(p)
        conn = connect(p)
        try:
            payload = execute_analysis_network_order_save(
                conn,
                position_key=position_key,
                order_raw=data.get("order"),
            )
            return jsonify(payload)
        except AnalysisUseCaseHTTP as e:
            return jsonify(e.payload), int(e.status_code)
        finally:
            conn.close()

    @app.post("/api/analysis/network-summary/export")
    @_perf_route("api/analysis/network-summary/export")
    @login_required
    def api_analysis_network_summary_export():
        """Deprecated: используйте POST /api/analysis/network-summary/export-job."""
        return (
            jsonify(
                {
                    "ok": False,
                    "error": "Экспорт переведён в фон. Используйте /api/analysis/network-summary/export-job",
                    "use_job": True,
                }
            ),
            410,
        )

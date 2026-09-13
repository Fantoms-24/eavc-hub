"""Импорт XLSX в таблицу online_search."""

from __future__ import annotations

from collections.abc import Callable
import logging
import re
import sqlite3
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from web_portal.config import DEFAULT_DB_NAME, db_path
from web_portal.lib.db import (
    connect,
    enqueue_sync_outbox,
    get_online_search_meta_payload,
    set_online_search_headers,
    sync_units_from_online_search,
)

from web_portal.application.online_search.errors import OnlineSearchUseCaseHTTP

_log = logging.getLogger("web_portal.application.online_search.import_xlsx")


def _norm_cell(v: object) -> str:
    s = str(v or "").strip().lower()
    s = s.replace("\n", " ").replace("\r", " ")
    s = re.sub(r"\s+", " ", s)
    return s


def _is_freq_header(s: str) -> bool:
    return ("частот" in s) or (s in {"freq", "frequency", "частота"})


def _is_gid_header(s: str) -> bool:
    return ("груп" in s and "id" in s) or s in {
        "group id",
        "group_id",
        "gid",
        "группа",
    }


def _is_note_header(s: str) -> bool:
    return any(
        x in s
        for x in (
            "примеч",
            "коммент",
            "comment",
            "note",
            "подраздел",
            "наимен",
            "назван",
            "unit",
        )
    )


def _is_id_header(s: str) -> bool:
    return ("id" in s) and ("груп" not in s) and ("group" not in s)


def _is_corr_header(s: str) -> bool:
    return ("корр" in s) or ("correspond" in s)


def execute_online_search_import_xlsx(
    conn: sqlite3.Connection,
    *,
    xlsx_path: Path,
    apply_to_unit: bool,
    import_id: str,
    user_id: int,
    is_sync_hub: bool,
    progress_set: Callable[[str, dict[str, Any]], None],
) -> dict[str, Any]:
    imported = 0
    unit_updated: dict[str, int] | None = {"inserted": 0, "updated": 0, "matched_pairs": 0}
    progress_set(
        import_id,
        {
            "user_id": user_id,
            "phase": "upload",
            "message": "Подготовка к импорту",
            "percent": 0,
            "processed_rows": 0,
            "total_rows": 0,
            "imported": 0,
            "done": False,
            "error": "",
        },
    )

    try:
        wb = load_workbook(str(xlsx_path), data_only=True)
        ws = wb.active

        selected_ws = ws
        header_row = 1
        freq_idx = 0
        gid_idx = 1
        note_idx = 8
        id_idx = 2
        corr_idx = 3
        headers = [""] * 9
        best_score = -1

        for ws_candidate in wb.worksheets:
            max_scan_row = min(30, int(ws_candidate.max_row or 0))
            for r in range(1, max_scan_row + 1):
                raw = [c.value for c in ws_candidate[r][:9]]
                vals = [_norm_cell(v) for v in raw]
                if not any(vals):
                    continue
                has_freq = any(_is_freq_header(v) for v in vals)
                has_gid = any(_is_gid_header(v) for v in vals)
                has_note = any(_is_note_header(v) for v in vals)
                has_id = any(_is_id_header(v) for v in vals)
                has_corr = any(_is_corr_header(v) for v in vals)
                score = (
                    (5 if has_freq else 0)
                    + (5 if has_gid else 0)
                    + (2 if has_note else 0)
                    + (1 if has_id else 0)
                    + (1 if has_corr else 0)
                )
                if score > best_score:
                    best_score = score
                    selected_ws = ws_candidate
                    header_row = r
                    headers = [str(v or "") for v in raw]
                    for i, v in enumerate(vals):
                        if _is_freq_header(v):
                            freq_idx = i
                        if _is_gid_header(v):
                            gid_idx = i
                        if _is_note_header(v):
                            note_idx = i
                        if _is_id_header(v):
                            id_idx = i
                        if _is_corr_header(v):
                            corr_idx = i
            if best_score >= 10:
                break

        ws = selected_ws
        total_rows = max(0, int(ws.max_row or 0) - int(header_row or 0))
        processed_rows = 0
        progress_set(
            import_id,
            {
                "phase": "parse",
                "message": "Парсинг строк",
                "total_rows": total_rows,
                "processed_rows": 0,
                "percent": 0,
                "imported": 0,
            },
        )

        set_online_search_headers(
            conn,
            headers,
            freq_col=freq_idx + 1,
            gid_col=gid_idx + 1,
            note_col=note_idx + 1,
        )
        conn.execute("DELETE FROM online_search")
        conn.commit()
        existing_map: dict[tuple[str, str, str, str], str] = {}

        if is_sync_hub:
            try:
                meta_payload = get_online_search_meta_payload(conn)
                main_conn = connect(db_path(DEFAULT_DB_NAME))
                try:
                    enqueue_sync_outbox(
                        main_conn, kind="online_search:meta", payload=meta_payload
                    )
                finally:
                    main_conn.close()
            except Exception:
                _log.debug("execute_online_search_import_xlsx: suppressed error", exc_info=True)

        main_conn = connect(db_path(DEFAULT_DB_NAME)) if apply_to_unit else None
        conn.execute("BEGIN")
        to_update: list[tuple[str, str, str, str, str]] = []
        to_insert: list[tuple[str, ...]] = []
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            processed_rows += 1
            if processed_rows % 200 == 0:
                percent = int((processed_rows * 100) / total_rows) if total_rows else 0
                progress_set(
                    import_id,
                    {
                        "phase": "parse",
                        "message": "Парсинг строк",
                        "processed_rows": processed_rows,
                        "total_rows": total_rows,
                        "percent": min(99, max(0, percent)),
                        "imported": imported,
                    },
                )
            vals = list(row[:9]) + [None] * (9 - len(row[:9]))
            if not any(v is not None and str(v).strip() != "" for v in vals):
                continue
            freq = str(vals[freq_idx] or "").strip()
            gid = str(vals[gid_idx] or "").strip()
            note = str(vals[note_idx] or "").strip()
            col3 = str(vals[id_idx] or "").strip() if 0 <= id_idx < len(vals) else ""
            col4 = str(vals[corr_idx] or "").strip() if 0 <= corr_idx < len(vals) else ""
            key = (freq, gid, col3, col4)
            prev_note = existing_map.get(key)
            if prev_note is not None:
                if prev_note == note:
                    continue
                to_update.append((note, freq, gid, col3, col4))
            else:
                cols = ["" if vals[i] is None else str(vals[i]) for i in range(9)]
                to_insert.append(
                    (
                        cols[0],
                        cols[1],
                        cols[2],
                        cols[3],
                        cols[4],
                        cols[5],
                        cols[6],
                        cols[7],
                        cols[8],
                        freq,
                        gid,
                        note,
                    )
                )
                existing_map[key] = note

        if to_update:
            conn.executemany(
                """
                UPDATE online_search
                SET note=?, updated_at=CURRENT_TIMESTAMP
                WHERE frequency=? AND group_id=? AND col3=? AND col4=?
                """,
                to_update,
            )
        if to_insert:
            conn.executemany(
                """
                INSERT INTO online_search
                  (col1,col2,col3,col4,col5,col6,col7,col8,col9, frequency, group_id, note)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?)
                """,
                to_insert,
            )

        imported = len(to_update) + len(to_insert)
        progress_set(
            import_id,
            {
                "phase": "write",
                "message": "Сохранение в БД",
                "processed_rows": processed_rows,
                "total_rows": total_rows,
                "percent": 99 if total_rows else 90,
                "imported": imported,
            },
        )

        wb.close()
        conn.commit()
        if main_conn:
            try:
                progress_set(
                    import_id,
                    {
                        "phase": "sync_unit",
                        "message": "Синхронизация подразделений",
                        "percent": 99,
                        "imported": imported,
                    },
                )
                unit_updated = sync_units_from_online_search(main_conn, conn)
            finally:
                main_conn.close()
        else:
            unit_updated = None

        progress_set(
            import_id,
            {
                "phase": "done",
                "message": "Импорт завершен",
                "processed_rows": processed_rows,
                "total_rows": total_rows,
                "percent": 100,
                "imported": imported,
                "done": True,
                "error": "",
            },
        )
        return {
            "ok": True,
            "imported": imported,
            "import_id": import_id,
            "unit_sync": unit_updated if apply_to_unit else None,
        }
    except OnlineSearchUseCaseHTTP:
        raise
    except Exception as e:
        progress_set(
            import_id,
            {
                "phase": "error",
                "message": "Ошибка импорта",
                "done": True,
                "error": str(e),
            },
        )
        raise OnlineSearchUseCaseHTTP(
            500, {"ok": False, "error": str(e), "import_id": import_id}
        ) from e

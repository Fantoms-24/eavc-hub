"""Экспорт сеансов из папки result0.txt → Excel (без Flask)."""

from __future__ import annotations

import logging
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Iterator

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from web_portal.application.admin.errors import AdminUseCaseHTTP

_log = logging.getLogger("web_portal.application.admin.export_sessions_excel")


def resolve_export_sessions_folder(folder_path: str) -> Path:
    path_s = (folder_path or "").strip()
    if not path_s:
        raise AdminUseCaseHTTP(
            400,
            {"ok": False, "error": "Укажите путь к папке (folder_path)"},
        )
    path = Path(path_s)
    if not path.is_dir():
        raise AdminUseCaseHTTP(
            400,
            {"ok": False, "error": f"Папка не найдена: {path}"},
        )
    return path


def iter_export_sessions_excel_events(folder: Path) -> Iterator[dict[str, Any]]:
    """
    Генератор событий прогресса экспорта.
    Успешное завершение: stage=done + content (bytes) + filename.
    Ошибка: stage=error + error (str).
    """
    from web_portal.lib.collect_seanses_from_dir import collect_seanses_from

    try:
        yield {
            "stage": "scan",
            "message": "Обработка файлов (без предварительного сканирования)...",
        }
        rows: list[tuple[Any, ...]] = []
        for idx, seans_list in enumerate(collect_seanses_from(folder)):
            yield {
                "stage": "file",
                "current": idx + 1,
                "message": f"Обработка файла {idx + 1}",
            }
            for e in seans_list:
                rows.append(
                    (
                        e.date_time,
                        e.frequency,
                        e.id,
                        e.group,
                        e.aes_key,
                        getattr(e, "color_voice", None),
                        getattr(e, "time_seconds", None),
                    )
                )

        if not rows:
            yield {
                "stage": "error",
                "error": "В папке не найдено файлов result0.txt или в них нет сеансов",
            }
            return

        yield {"stage": "excel", "message": "Формирование Excel..."}
        rows.sort(key=lambda r: (r[0], r[2]))

        wb = Workbook()
        ws = wb.active
        ws.title = "Сеансы"
        headers = [
            "Время выхода",
            "Частота",
            "ID корреспондента",
            "Группа",
            "AES ключа",
            "Color Voice",
            "Время выхода (сек)",
        ]
        for col, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        for row_idx, r in enumerate(rows, start=2):
            ws.cell(row=row_idx, column=1, value=r[0])
            ws.cell(row=row_idx, column=2, value=r[1])
            ws.cell(row=row_idx, column=3, value=r[2])
            ws.cell(row=row_idx, column=4, value=r[3])
            ws.cell(row=row_idx, column=5, value=r[4] or "")
            ws.cell(row=row_idx, column=6, value=r[5] or "")
            ws.cell(
                row=row_idx,
                column=7,
                value=round(r[6], 3) if r[6] is not None else "",
            )
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 18

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f"seanses_export_{datetime.now().strftime('%Y-%m-%d_%H-%M-%S')}.xlsx"
        yield {
            "stage": "done",
            "filename": filename,
            "content": buf.getvalue(),
        }
    except Exception as e:
        _log.exception("export-sessions-excel: %s", e)
        yield {"stage": "error", "error": str(e)}
